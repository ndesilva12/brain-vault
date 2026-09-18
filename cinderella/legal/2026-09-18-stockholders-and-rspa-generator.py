# -*- coding: utf-8 -*-
"""Stockholders' Agreement + Restricted Stock Purchase Agreement for Cinderella Corp."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = ("/tmp/claude-0/-home-user-brain-vault/cdeb7b3c-9b38-5f86-b0fa-23ce6e519c4c/scratchpad/")


def newdoc():
    d = Document()
    for s in d.sections:
        s.top_margin = s.bottom_margin = Inches(1.0)
        s.left_margin = s.right_margin = Inches(1.0)
    st = d.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(10.5)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_after = Pt(9)
    st.paragraph_format.line_spacing = 1.06
    return d


def mk(d):
    def p(text="", *, lead=None, indent=0.0, center=False, italic=False, bold=False,
          before=0, after=9, size=10.5, justify=True):
        par = d.add_paragraph(); pf = par.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        if indent: pf.left_indent = Inches(indent)
        if center: par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif justify: par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if lead:
            r = par.add_run(lead); r.bold = True; r.font.size = Pt(size)
        if text:
            r = par.add_run(text); r.italic = italic; r.bold = bold; r.font.size = Pt(size)
        return par

    def head(t, size=11.5, before=15, after=7):
        par = d.add_paragraph(); pf = par.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        r = par.add_run(t); r.bold = True; r.font.size = Pt(size)

    def title(t, size=13):
        par = d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(14)
        r = par.add_run(t); r.bold = True; r.font.size = Pt(size)

    def sig(label, extra=None):
        par = d.add_paragraph(); par.paragraph_format.space_before = Pt(15)
        par.paragraph_format.space_after = Pt(2); par.add_run("_" * 44)
        if label:
            q = d.add_paragraph(); q.paragraph_format.space_after = Pt(0); q.add_run(label)
        for e in (extra or []):
            q = d.add_paragraph(); q.paragraph_format.space_after = Pt(0); q.add_run(e)
    return p, head, title, sig


# ═══════════════════════════ STOCKHOLDERS' AGREEMENT ═══════════════════════════
d = newdoc(); p, head, title, sig = mk(d)

p("DRAFT — FOR DISCUSSION ONLY — NOT LEGAL ADVICE", bold=True, center=True, size=9,
  after=12, justify=False)
title("STOCKHOLDERS’ AGREEMENT")
p("Cinderella Corp., a Delaware corporation", center=True, italic=True, size=10, after=14,
  justify=False)

p("This Stockholders’ Agreement (this “Agreement”) is entered into as of [____], 2026 "
  "(the “Effective Date”) by and among Cinderella Corp., a Delaware corporation (the "
  "“Company”), Norman C. de Silva (the “Founder”), and each other holder of capital "
  "stock of the Company who executes this Agreement or a joinder hereto (each, together with the "
  "Founder, a “Stockholder”).")

head("1.  CAPITAL STRUCTURE")
p("The Company’s authorized capital stock consists of 6,000,000 shares of Class A Common "
  "Stock and 4,000,000 shares of Class B Common Stock, each with a par value of $0.001.",
  lead="1.1  Authorized capital.  ")
p("Each share of Class A Common Stock carries one (1) vote. Each share of Class B Common Stock "
  "carries ten (10) votes. Except as otherwise required by law or the Certificate of "
  "Incorporation, the classes vote together as a single class. Except as to voting and "
  "conversion, the classes are identical and carry equal per-share rights to dividends and to "
  "distributions upon liquidation.", lead="1.2  Voting.  ")
p("The Founder holds all outstanding Class B Common Stock. No person other than the Founder holds "
  "or may hold Class B Common Stock. All other Stockholders hold Class A Common Stock only.",
  lead="1.3  Holdings.  ")
p("Each share of Class B Common Stock converts automatically into one share of Class A Common "
  "Stock upon any transfer other than a Permitted Transfer, and converted shares are retired and "
  "may not be reissued as Class B Common Stock.", lead="1.4  Automatic conversion.  ")

head("2.  GOVERNANCE")
p("The Board of Directors consists of three (3) authorized seats. As of the Effective Date, the "
  "Founder occupies one seat and two seats are vacant. Directors are elected by the vote of the "
  "holders of the Company’s outstanding voting stock.", lead="2.1  Board of Directors.  ")
p("The Company is managed under the direction of the Board. Except as expressly provided in "
  "Section 2.3, no Stockholder other than the Founder has any right to designate a director, to "
  "attend meetings of the Board as an observer, or to participate in the management of the "
  "Company.", lead="2.2  Management.  ")
p("Ankur Jain (the “Investor Representative”) may attend all meetings of the Board in a "
  "non-voting observer capacity and shall receive copies of all materials provided to directors, "
  "at the same time they are provided. Upon the appointment or election of any additional "
  "director to the Board — that is, at such time as any person other than the Founder becomes a "
  "director — the Investor Representative shall automatically and concurrently be appointed to "
  "the Board as a director, without further action or discretion on the part of the Company or "
  "the Founder, and each Stockholder shall vote all shares in favor of that appointment. The "
  "Investor Representative may decline the seat and elect to continue as an observer, in which "
  "case the right to take the seat continues and may be exercised at any later time on written "
  "notice.", lead="2.3  Board observer; automatic seat.  ")
p("The Company may exclude the observer from any portion of a meeting, and may withhold any "
  "material, where the Board determines in good faith that doing so is necessary to preserve "
  "attorney-client privilege, to address an actual or potential conflict of interest involving "
  "the observer or any Stockholder he represents, or to protect highly confidential information. "
  "The observer is bound by Section 7.", lead="2.4  Observer exclusions.  ")
p("The Investor Representative is authorized to act on behalf of the Investors with respect to "
  "the matters expressly assigned to him under this Agreement and under the Side Letter "
  "Agreement. The Company is entitled to rely conclusively on any action, notice, consent or "
  "waiver given by the Investor Representative as binding on all Investors. The Investor "
  "Representative may be replaced by Investors holding a majority of the shares held by Investors "
  "upon written notice to the Company.", lead="2.5  Investor Representative.  ")
p("The Company will indemnify, defend and hold harmless the Investor Representative, and advance "
  "his reasonable expenses, to the fullest extent permitted by Delaware law, in respect of any "
  "claim arising out of his good-faith performance of the role, including any claim brought by an "
  "Investor. No indemnity is available for bad faith, willful misconduct, gross negligence, or "
  "breach of this Agreement, and advancement is conditioned on an undertaking to repay amounts "
  "advanced if it is ultimately determined that he was not entitled to indemnification. The "
  "Company will use commercially reasonable efforts to include the Investor Representative under "
  "its directors’ and officers’ liability insurance in his capacities as Board "
  "observer or director and as Investor Representative.",
  lead="2.6  Indemnification of the Investor Representative.  ")

head("3.  TRANSFER RESTRICTIONS")
p("No Stockholder may transfer any shares without (a) the prior written consent of the Board and "
  "(b) compliance with the Securities Act of 1933, as amended, and applicable state securities "
  "laws. Any purported transfer in violation of this Article is void.",
  lead="3.1  General restriction.  ")
p("The following are “Permitted Transfers” and do not require Board consent, provided "
  "the transferee executes a joinder to this Agreement: a transfer to a spouse, lineal descendant "
  "or other family member; to a trust or other vehicle established for estate-planning purposes; "
  "or by will or intestacy. Permitted Transfers of Class B Common Stock by the Founder are "
  "governed by the Certificate of Incorporation.", lead="3.2  Permitted Transfers.  ")
p("Before transferring any shares to a third party, a Stockholder must first offer them to the "
  "Company, and if the Company declines or does not respond within thirty (30) days, to the other "
  "Stockholders pro rata, in each case on the same terms and for a period of thirty (30) days.",
  lead="3.3  Right of first refusal.  ")
p("If the Founder proposes to transfer shares to a third party in a transaction that is not a "
  "Permitted Transfer, each other Stockholder may elect to include in that transfer a pro rata "
  "portion of that Stockholder’s shares on the same terms.", lead="3.4  Co-sale.  ")

head("4.  DRAG-ALONG")
p("If the Board approves a sale of the Company — whether structured as a merger, consolidation, "
  "share exchange, sale of stock, or sale of all or substantially all of the assets of the "
  "Company — each Stockholder will (a) vote all shares in favor of the transaction and against "
  "any competing proposal, (b) participate on the same per-share terms applicable to holders of "
  "the same class, (c) refrain from exercising any appraisal or dissenters’ rights, (d) "
  "execute the documents reasonably necessary to effect the transaction, and (e) not be required "
  "to make representations other than as to title to the shares, authority, and the absence of "
  "liens, or to bear liability beyond that Stockholder’s pro rata share of the "
  "consideration received. This Article survives until the closing of such a sale.",
  lead="4.1  Obligation.  ")

head("5.  REPURCHASE RIGHTS")
p("Shares issued subject to vesting are governed by the applicable Restricted Stock Purchase "
  "Agreement, which provides for repurchase of unvested shares at the lower of cost or fair "
  "market value on termination of service.", lead="5.1  Unvested shares.  ")
p("Nothing in this Agreement limits any repurchase, forfeiture or clawback right set forth in a "
  "Stockholder’s individual agreement with the Company.", lead="5.2  Other rights preserved.  ")

head("6.  INFORMATION RIGHTS")
p("The Company will furnish to each Stockholder holding at least one percent (1%) of the "
  "outstanding capital stock: (a) unaudited annual financial statements within one hundred twenty "
  "(120) days after each fiscal year end; and (b) material financing updates as they occur. "
  "Information rights terminate on the closing of a sale of the Company.",
  lead="6.1  Reporting.  ")
p("The rights in Section 6.1 are a floor and not a ceiling. Where a Stockholder’s separate "
  "written agreement with the Company provides for additional or more frequent information — "
  "including quarterly financial statements and the annual budget — that agreement governs as to "
  "that Stockholder.", lead="6.2  Additional rights preserved.  ")
p("Nothing in this Article entitles any Stockholder to the Company’s trade secrets, the "
  "Format or franchise intellectual property, privileged materials, or information the Board "
  "determines in good faith to be competitively sensitive.", lead="6.3  Limits.  ")

head("7.  CONFIDENTIALITY AND NON-CIRCUMVENTION")
p("Each Stockholder will hold in confidence all non-public information regarding the Company and "
  "use it solely for purposes of evaluating and monitoring that Stockholder’s investment or "
  "performing that Stockholder’s role. This obligation does not apply to information that is "
  "or becomes public other than through that Stockholder’s breach, or that the Stockholder "
  "is required to disclose by law, provided the Stockholder gives the Company prompt notice where "
  "lawful to do so.", lead="7.1  Confidentiality.  ")
p("During the term and for twelve (12) months thereafter, no Stockholder will, using the "
  "Company’s confidential information or an opportunity originated by the Company, "
  "circumvent the Company with respect to any project that is substantially similar to the "
  "Company’s business of pairing celebrities or public figures with collegiate athletic "
  "programs and producing documentary content concerning them. Nothing in this Section restricts "
  "any Stockholder from initiating, continuing or maintaining any personal or professional "
  "relationship with any person, whenever and however that relationship arose.",
  lead="7.2  Non-circumvention.  ")
p("Where a Stockholder’s separate written agreement with the Company contains a "
  "non-circumvention or non-solicitation provision, including any definition of a restricted "
  "project and any carve-outs, that provision governs as to that Stockholder and Section 7.2 does "
  "not operate to broaden it.", lead="7.3  Separate agreements control.  ")

head("8.  LEGENDS")
p("All certificates and book-entry positions representing shares subject to this Agreement bear a "
  "legend referencing these restrictions and the restrictions of the Securities Act.",
  lead="8.1  Legend.  ")

head("9.  GENERAL")
p("Delaware law governs, without regard to conflicts of law principles. The Court of Chancery of "
  "the State of Delaware has exclusive jurisdiction. TO THE FULLEST EXTENT PERMITTED BY LAW, EACH "
  "PARTY WAIVES ANY RIGHT TO TRIAL BY JURY.", lead="9.1  Governing law.  ")
p("This Agreement may be amended by the Company with the written consent of the holders of a "
  "majority of the outstanding Class B Common Stock and the holders of a majority of the "
  "outstanding Class A Common Stock, and any such amendment binds all Stockholders. Any amendment "
  "that disproportionately and adversely affects the Investors relative to other holders of Class "
  "A Common Stock also requires the consent of the Investor Representative.",
  lead="9.2  Amendment.  ")
p("This Agreement terminates upon the closing of a sale of the Company or upon the written "
  "agreement of the Company and all Stockholders.", lead="9.3  Termination.  ")
p("In writing, to the addresses set forth on the signature pages or a joinder.",
  lead="9.4  Notices.  ")
p("If any provision is held invalid or unenforceable, it is ineffective only to that extent and "
  "the remainder continues in full force.", lead="9.5  Severability.  ")
p("This Agreement may be executed in counterparts and by electronic signature. Any person may "
  "become a party by executing a counterpart signature page or joinder.",
  lead="9.6  Counterparts; joinder.  ")
p("This Agreement sets out rights and obligations common to all Stockholders. Where a "
  "Stockholder’s separate written agreement with the Company addresses the same subject "
  "matter — including vesting, acceleration, forfeiture, repurchase, information rights, "
  "governance participation, indemnification or non-circumvention — that separate agreement "
  "controls as to that Stockholder, and nothing in this Agreement reduces or conditions a right "
  "granted there.", lead="9.7  Separate agreements.  ")

p("[Remainder of page intentionally left blank]", center=True, italic=True, before=14, after=14,
  justify=False)
d.add_page_break()
p("IN WITNESS WHEREOF, the parties have executed this Stockholders’ Agreement as of the date "
  "first above written.", after=16)
p("COMPANY:", bold=True, after=0, justify=False)
p("CINDERELLA CORP.", before=8, after=0, justify=False)
sig("By:", ["Name:  Norman C. de Silva", "Title:  Founder & Chief Executive Officer"])
p("FOUNDER:", bold=True, before=20, after=0, justify=False)
sig("Norman C. de Silva")
p("STOCKHOLDER:", bold=True, before=20, after=0, justify=False)
sig("Name:", ["Address for notices:  ______________________________", "Email:  ______________________"])

d.save(BASE + "Cinderella Corp - Stockholders Agreement (DRAFT).docx")

DEFAULT_SCHEDULE = [
    ("Number of Shares:  ______________", dict(after=4, justify=False)),
    ("Purchase price per Share:  $______________", dict(after=4, justify=False)),
    ("Vesting commencement date:  ______________", dict(after=12, justify=False)),
    ("Vesting schedule:", dict(bold=True, after=6, justify=False)),
    ("[Insert the schedule from the holder’s separate agreement. Examples:]",
     dict(italic=True, after=6)),
    ("\u2022  Fully vested on the Effective Date; no repurchase option applies except as provided "
     "in the holder’s separate agreement.", dict(indent=0.25, after=4)),
    ("\u2022  Fifty percent (50%) vested on the Effective Date; the remainder in twenty-four (24) "
     "equal monthly installments, vesting on the last day of each full calendar month thereafter, "
     "subject to continuous service.", dict(indent=0.25, after=4)),
    ("\u2022  Monthly over thirty-six (36) months from the vesting commencement date, no cliff, "
     "subject to continuous service.", dict(indent=0.25, after=4)),
    ("\u2022  Milestone grant under Section 2.4. Issued on the Company’s written determination "
     "that the [____] milestone under the holder’s separate agreement dated [____] has been "
     "achieved, and fully vested on issuance. No repurchase option applies.", dict(indent=0.25)),
]

# ═══════════════════ RESTRICTED STOCK PURCHASE AGREEMENT ═══════════════════
def build_rspa(outname, *, holder="[HOLDER NAME]", shares="[______]",
               price="$[____]", schedule=None, note=None):
    global d, p, head, title, sig
    d = newdoc(); p, head, title, sig = mk(d)

    p("DRAFT — FOR DISCUSSION ONLY — NOT LEGAL ADVICE", bold=True, center=True, size=9,
      after=12, justify=False)
    title("RESTRICTED STOCK PURCHASE AGREEMENT")
    p("Cinderella Corp., a Delaware corporation", center=True, italic=True, size=10, after=14,
      justify=False)

    p("This Restricted Stock Purchase Agreement (this “Agreement”) is entered into as of "
      "[____], 2026 (the “Effective Date”) by and between Cinderella Corp., a Delaware "
      "corporation (the “Company”), and " + holder + " (“Purchaser”).")
    if note:
        p(note, italic=True, size=9.5, after=12)

    head("1.  PURCHASE AND SALE")
    p("The Company issues and sells to Purchaser, and Purchaser purchases, " + shares +
      " shares of Class A Common Stock, par value $0.001 (the “Shares”), at a purchase "
      "price of " + price + " per Share.", lead="1.1  Shares.  ")
    p("Purchaser shall pay the aggregate purchase price in cash, by check, or — where the Board so "
      "determines and applicable law permits — in consideration of past or future services rendered "
      "to the Company, in each case as recorded in the Company’s books.",
      lead="1.2  Consideration.  ")
    p("The Shares are issued pursuant to the exemption from registration provided by [Rule 701 / "
      "Regulation D] under the Securities Act of 1933, as amended. The Shares have not been "
      "registered and may not be transferred absent registration or an available exemption.",
      lead="1.3  Securities exemption.  ")

    head("2.  VESTING")
    p("The Shares vest in accordance with Schedule A. Shares that have vested are "
      "“Vested Shares”; all others are “Unvested Shares.”",
      lead="2.1  Schedule.  ")
    p("Except as expressly provided in Schedule A or in Purchaser’s separate written agreement "
      "with the Company, vesting is subject to Purchaser’s continuous service with the Company "
      "through each applicable vesting date.", lead="2.2  Service condition.  ")
    p("Any acceleration of vesting provided in Purchaser’s separate written agreement with the "
      "Company is incorporated by reference and applies to the Shares.",
      lead="2.3  Acceleration.  ")
    p("Where Schedule A identifies the Shares as a milestone grant, the Shares are issued only upon "
      "the Company’s written determination that the applicable milestone under Purchaser’s "
      "separate written agreement has been achieved, and are fully vested on issuance. Article 3 does "
      "not apply to such Shares. The milestone itself, its definition, and any cap on the aggregate "
      "equity issuable across milestones are governed exclusively by that separate agreement.",
      lead="2.4  Milestone grants.  ")

    head("3.  REPURCHASE OPTION")
    p("On termination of Purchaser’s service with the Company for any reason, the Company has "
      "the option, exercisable for ninety (90) days following termination, to repurchase all or any "
      "portion of the Unvested Shares at the lower of (a) the price paid for them and (b) their fair "
      "market value as determined in good faith by the Board.",
      lead="3.1  Unvested Shares.  ")
    p("The Company may exercise the option by written notice, and may assign the option to any person "
      "it designates. On exercise, the Shares are deemed repurchased as of the notice date and "
      "Purchaser ceases to have any rights in them other than the right to receive the repurchase "
      "price.", lead="3.2  Mechanics.  ")
    p("Any repurchase, forfeiture or clawback right set forth in Purchaser’s separate written "
      "agreement with the Company applies in addition to this Article.",
      lead="3.3  Other rights preserved.  ")

    head("4.  TRANSFER RESTRICTIONS")
    p("The Shares are subject to the Stockholders’ Agreement, including its transfer "
      "restrictions, right of first refusal, co-sale and drag-along provisions. Purchaser is a party "
      "to the Stockholders’ Agreement or shall execute a joinder to it concurrently with this "
      "Agreement.", lead="4.1  Stockholders’ Agreement.  ")
    p("Unvested Shares may not be transferred under any circumstance other than to the Company.",
      lead="4.2  Unvested Shares.  ")

    head("5.  TAX")
    p("Purchaser is strongly advised to consult his or her own tax adviser and to consider filing an "
      "election under Section 83(b) of the Internal Revenue Code within thirty (30) days of the "
      "Effective Date. The election is Purchaser’s sole responsibility; the Company makes no tax "
      "representation and has no obligation to file it on Purchaser’s behalf. A form of election "
      "is attached as Schedule B.", lead="5.1  83(b) election.  ")
    p("The Company may withhold, or require payment of, applicable taxes as a condition to issuance "
      "or vesting.", lead="5.2  Withholding.  ")

    head("6.  MARKET STAND-OFF")
    p("If requested by the managing underwriter of an initial public offering of the Company, "
      "Purchaser will not sell, transfer or otherwise dispose of any securities of the Company for "
      "such period following the offering as the underwriter requests, not to exceed one hundred "
      "eighty (180) days and no longer than the shortest period required of any other holder.",
      lead="6.1  Lock-up.  ")

    head("7.  LEGENDS")
    p("Certificates or book-entry positions representing the Shares bear legends referencing the "
      "restrictions of the Securities Act, this Agreement, and the Stockholders’ Agreement.",
      lead="7.1  Legends.  ")

    head("8.  REPRESENTATIONS OF PURCHASER")
    p("Purchaser represents that: (a) Purchaser is acquiring the Shares for Purchaser’s own "
      "account, for investment, and not with a view to distribution; (b) Purchaser can bear the "
      "economic risk of the investment, including its total loss; (c) Purchaser has had the "
      "opportunity to ask questions of the Company and to consult independent legal and tax counsel; "
      "and (d) Purchaser understands the Shares are illiquid, subject to vesting and repurchase, and "
      "may never have any value.", lead="8.1  Representations.  ")

    head("9.  GENERAL")
    p("Delaware law governs, without regard to conflicts of law principles; the Court of Chancery of "
      "the State of Delaware has exclusive jurisdiction.", lead="9.1  Governing law.  ")
    p("This Agreement, together with the Stockholders’ Agreement and Purchaser’s separate "
      "written agreement with the Company, constitutes the entire agreement as to its subject matter. "
      "In the event of a conflict as to vesting, acceleration, forfeiture or repurchase, "
      "Purchaser’s separate written agreement controls.", lead="9.2  Entire agreement.  ")
    p("Amendments in writing signed by both parties. Purchaser may not assign. Counterparts and "
      "electronic signature permitted. If any provision is held unenforceable, the remainder "
      "continues in full force.", lead="9.3  Miscellaneous.  ")

    p("[Remainder of page intentionally left blank]", center=True, italic=True, before=14, after=14,
      justify=False)
    d.add_page_break()
    p("IN WITNESS WHEREOF, the parties have executed this Restricted Stock Purchase Agreement as of "
      "the date first above written.", after=16)
    p("COMPANY:", bold=True, after=0, justify=False)
    p("CINDERELLA CORP.", before=8, after=0, justify=False)
    sig("By:", ["Name:  Norman C. de Silva", "Title:  Founder & Chief Executive Officer"])
    p("PURCHASER:", bold=True, before=20, after=0, justify=False)
    sig("Name:", ["Address:  ______________________________", "Email:  ______________________"])

    d.add_page_break()
    title("SCHEDULE A — VESTING", size=12)
    for _line in (schedule or DEFAULT_SCHEDULE):
        _txt, _kw = _line
        p(_txt, **_kw)

    d.add_page_break()
    title("SCHEDULE B — FORM OF 83(b) ELECTION", size=12)
    p("[Attach the standard Internal Revenue Code Section 83(b) election form. Purchaser must file it "
      "with the Internal Revenue Service within thirty (30) days of the Effective Date and provide a "
      "copy to the Company.]", italic=True)

    d.save(BASE + outname)



build_rspa("Cinderella Corp - Restricted Stock Purchase Agreement (DRAFT).docx")

build_rspa(
    "Cinderella Corp - RSPA - Ankur Jain (DRAFT).docx",
    holder="Ankur Jain",
    shares="166,667",
    price="$[____]",
    note=("Completed for Ankur Jain pursuant to the Strategic Adviser Agreement dated "
          "[____], 2026. Share count reflects 4% of the Company’s fully-diluted "
          "capitalization as of the Effective Date, against 3,000,000 shares of Class B Common "
          "Stock held by the Founder and a pre-Seed fully-diluted total of 4,166,667 shares. "
          "Purchase price per Share to be set at fair market value as determined in good faith "
          "by the Board — confirm with counsel before execution."),
    schedule=[
        ("Purchaser:  Ankur Jain", dict(after=4, justify=False)),
        ("Number of Shares:  166,667 shares of Class A Common Stock", dict(after=4, justify=False)),
        ("Purchase price per Share:  $______________  [FMV as determined by the Board]",
         dict(after=4, justify=False)),
        ("Percentage of the Company:  4.000% of the fully-diluted capitalization as of the "
         "Effective Date (pre-Seed fully-diluted total: 4,166,667 shares)",
         dict(after=4, justify=False)),
        ("Vesting commencement date:  Not applicable", dict(after=12, justify=False)),
        ("Vesting schedule:", dict(bold=True, after=6, justify=False)),
        ("\u2022  Fully vested on the Effective Date. No repurchase option applies, and Article 3 "
         "of this Agreement does not apply to the Shares, per Section 2.1 of the Strategic "
         "Adviser Agreement.", dict(indent=0.25, after=10)),
        ("Seed Round top-up:", dict(bold=True, after=6, justify=False)),
        ("\u2022  Under Section 2.4 of the Strategic Adviser Agreement, at the closing of the Seed "
         "Round the Company will issue Purchaser, for no additional consideration, such "
         "additional shares of Class A Common Stock as are necessary for the Shares to represent "
         "4% of the fully-diluted capitalization immediately following that closing. Based on "
         "current capitalization the top-up is expected to be approximately 30,864 shares, for a "
         "post-Seed total of approximately 197,531 shares of a 4,938,272-share fully-diluted "
         "total. The top-up shares are fully vested on issuance and are treated as Shares under "
         "this Agreement. Actual numbers are those required to produce 4% and are subject to "
         "rounding.", dict(indent=0.25, after=4)),
    ])
