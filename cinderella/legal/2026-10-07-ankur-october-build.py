# -*- coding: utf-8 -*-
"""Apply Ankur's 2026-10-07 redline to the five current documents and emit .docx.

SOURCE OF TRUTH: the text extracted from the PDFs Norman sent 2026-10-07, in
legal/current-versions-2026-10-07/. This is a PATCH build, not a re-draft — every change below
is a targeted replacement against that text, and each one is asserted to have fired. If a source
string changes, the build fails loudly rather than silently skipping a change.

NORMAN'S DECISIONS (2026-10-07):
  • Securities exemption → Regulation D / Section 4(a)(2) + Rule 506(b). Rule 701 rejected
    because 701 excludes services rendered in connection with capital raising, and Adviser §1.1
    defines the services as capital formation and investor meetings.
  • D&O insurance → option (b), bind at Seed Round close. So "commercially reasonable efforts
    within 90 days" becomes a firm trigger tied to the Seed closing, in all three places it
    appears.
"""
import re, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = Path(__file__).parent / "current-versions-2026-10-07"
OUT = Path(__file__).parent / "out-2026-10-07"
OUT.mkdir(exist_ok=True)

# ───────────────────────── text loading / normalisation ─────────────────────────

TABLE_MARK = "\x00PLATFORM_TABLE\x00"

def load(name):
    raw = (SRC / f"{name}.txt").read_text(encoding="utf-8")
    raw = raw.replace("\x0c", "\n\n")                      # form feeds -> para break
    raw = re.sub(r"[ \t]*\n[ \t]*\n[ \t]*(\n[ \t]*)+", "\n\n", raw)
    return raw

TABLE_HEADER = "Qualifying Platform    Meeting (25%)"

def excise_platform_table(text):
    """Pull the Qualifying Platform tier grid out so normalisation cannot mangle it.

    ⚠️ The start anchor MUST be the table's own header row, not the first occurrence of
    "Qualifying Platform". Ankur reported 2026-10-07 that his Adviser Agreement was missing the
    §2.2 heading and §2.2(a): this function used to anchor on text.find("Qualifying Platform"),
    whose first hit is inside §2.2's OWN opening paragraph ('...with respect to a Qualifying
    Platform. A "Qualifying Platform" means Netflix, Apple...'). So the excision swallowed the
    §2.2 heading, the definition of Qualifying Platform and all of §2.2(a) Platform tiers, and
    replaced them with the bare rebuilt grid — the document jumped from §2.1 straight to a table
    and then to §2.2(b). Anchoring on the header row leaves the prose intact.
    """
    if "(a) Platform tiers" not in text:
        return text, False          # only the Adviser Agreement carries the grid
    start = text.find(TABLE_HEADER)
    assert start != -1, ("this document has the platform-tier subsection but not the expected "
                         "table header row; a looser anchor would swallow §2.2's heading and "
                         "§2.2(a), which is exactly the bug Ankur reported on 2026-10-07")
    end = text.find("(b) Milestones", start)
    assert end != -1, "platform table end anchor '(b) Milestones' not found"
    return text[:start] + TABLE_MARK + "\n\n" + text[end:], True

def paragraphs(text):
    """Blank-line-delimited paragraphs; wrapped lines rejoined into one line each."""
    out = []
    for block in text.split("\n\n"):
        block = block.strip("\n")
        if not block.strip():
            continue
        if TABLE_MARK in block:
            # ⚠️ The table marker can share a block with prose, because §2.2(a) "Platform tiers"
            # runs straight into the table header row with no blank line between them. Discarding
            # the whole block loses §2.2(a) — half of the bug Ankur reported on 2026-10-07. Emit
            # the prose on either side of the marker instead.
            before, _, after = block.partition(TABLE_MARK)
            for side in (before,):
                lines = [ln.strip() for ln in side.split("\n") if ln.strip()]
                if lines:
                    out.append(" ".join(lines))
            out.append(TABLE_MARK)
            lines = [ln.strip() for ln in after.split("\n") if ln.strip()]
            if lines:
                out.append(" ".join(lines))
            continue
        lines = [ln.strip() for ln in block.split("\n") if ln.strip()]
        # a block of short all-caps lines is a heading cluster, keep them separate
        if all(len(ln) < 70 and ln == ln.upper() for ln in lines) and len(lines) > 1:
            out.extend(lines)
        # a field list (Purchaser: / Number of Shares: / Purchase price per Share:) must stay
        # one field per line, or a Schedule reads as a run-on paragraph
        elif sum(1 for ln in lines if re.match(r"^[A-Z][A-Za-z /\-]{2,40}:\s", ln)) >= 2:
            buf = []
            for ln in lines:
                if re.match(r"^[A-Z][A-Za-z /\-]{2,40}:\s", ln):
                    if buf:
                        out.append(" ".join(buf)); buf = []
                    out.append(ln)
                else:
                    if out and not re.match(r"^[A-Z][A-Za-z /\-]{2,40}:\s", ln):
                        out[-1] = out[-1] + " " + ln
                    else:
                        buf.append(ln)
            if buf:
                out.append(" ".join(buf))
        else:
            out.append(" ".join(lines))
    return out

# ───────────────────────── the change set ─────────────────────────
# (label, find, replace) — `find` matched against normalised paragraph text.
# A find of None means "insert_after the paragraph whose text starts with the anchor".

ADV_11 = ("Adviser serves as a non-exclusive Strategic Adviser, providing: (a) advice on financial "
          "structure, valuation, and capital formation; (b) strategic input on the Company's financial "
          "presentation as set forth in Section 1.2; (c) participation in PE/investor and distributor "
          "meetings (subject to OBA Approval and Section 1.5); and (d) introductions across his network.")

ADV_12 = ("1.2 Financial presentation — consultation; no approval or review duty. The Company will "
          "consult Adviser in good faith on the Company's financial presentation, valuation, financing "
          "structure and proposed investor terms, and Adviser may provide strategic input as he considers "
          "appropriate. Adviser has no duty to review, approve, verify, audit or monitor any financial "
          "materials, model, assumption, valuation or investor term, and no right of approval over any of "
          "them. Responsibility for the preparation, accuracy and completeness of the Company's financial "
          "materials rests with the Company, its Board of Directors and the accounting firm retained under "
          "Section 1.4. Adviser's participation in, or receipt of, any materials does not constitute "
          "approval, verification or endorsement of them, and his silence or non-response does not "
          "constitute approval.")

ADV_13 = ("1.3 No condition; no remedy. The consultation contemplated by Section 1.2 is a covenant of the "
          "Company only. It is not a condition precedent to any financing, issuance of securities, "
          "acceptance of any subscription, or any other corporate action, and no failure to consult shall "
          "invalidate, delay or give rise to any claim in respect of any such action. The issuance of "
          "securities, the consideration for which they are issued, and the acceptance of any subscription "
          "remain matters for determination by the Board under the Delaware General Corporation Law, the "
          "Company's Certificate of Incorporation and its Bylaws.")

ADV_25 = ("2.5 Anti-dilution — Seed Round only. Adviser's equity is measured on the fully-diluted "
          "capitalization immediately after completion of the Seed Round, as and when completed and "
          "whatever its final size, price and structure, so the Seed Round does not dilute Adviser. After "
          "the Seed Round is complete, Adviser's shares are subject to ordinary, pro-rata dilution on the "
          "same basis as the Founder on all subsequent issuances — no further anti-dilution protection. "
          "This post-Seed measurement is specific to Adviser and is not a precedent for any other grant.")

ADV_27 = ("2.7 Mechanics; tax. Issued as restricted stock under the Company's Restricted Stock Purchase "
          "Agreement and subject to the Stockholders' Agreement. The Role Shares and any milestone shares "
          "are fully vested on issuance and accordingly an election under Section 83(b) of the Internal "
          "Revenue Code is inapplicable to them; Adviser should consult his own tax advisor as to the "
          "treatment of each grant.")

ADV_43 = ("4.3 D&O insurance. The Company shall obtain directors' and officers' liability insurance on or "
          "before the closing of the Seed Round, and shall cause Adviser to be covered under that policy "
          "expressly in his capacities as Board observer or director and as Investor Representative. The "
          "Company shall provide Adviser with the policy's coverage limits and material terms upon request. "
          "The Company represents that, as of the Effective Date, it does not maintain such insurance.")

RSPA_11 = ("1.1 Shares. The Company issues and sells to Purchaser, and Purchaser purchases, 166,667 shares "
           "of Class A Common Stock, par value $0.001 (the “Shares”), at a purchase price of "
           "$0.001 per Share, which the Board has determined in good faith to be the fair market value of a "
           "share of Class A Common Stock as of the Effective Date.")

RSPA_13 = ("1.3 Securities exemption. The Shares are issued in a transaction exempt from registration under "
           "Section 4(a)(2) of the Securities Act of 1933, as amended, and Rule 506(b) of Regulation D "
           "thereunder. The Shares have not been registered and may not be transferred absent registration "
           "or an available exemption.")

RSPA_24 = ("2.4 Milestone grants. Where Schedule A identifies the Shares as a milestone grant, the Shares are "
           "issuable upon objective achievement of the applicable milestone as defined in Purchaser's "
           "separate written agreement with the Company, and are fully vested on issuance. Article 3 does "
           "not apply to such Shares. Issuance is not subject to any further determination, consent or "
           "discretion of the Company or the Board. Purchaser shall deliver written notice to the Company "
           "identifying the milestone achieved and the facts establishing achievement and attribution, and "
           "the Company shall issue the Shares within fifteen (15) business days following that notice. If "
           "the Company disputes achievement or attribution, it shall deliver written notice of the dispute, "
           "stating the grounds in reasonable detail, within ten (10) business days of Purchaser's notice; "
           "absent such timely notice the milestone is deemed achieved and attributable. The milestone "
           "definitions, the attribution requirement under Section 2.2(d) of that agreement, and the "
           "aggregate cap on equity issuable across milestones are governed exclusively by that agreement.")

# NOTE: curly apostrophes deliberately — this sentence must be character-identical to the same
# sentence in Sunjay's and Greg's copies, which use the document's curly typography throughout.
# The Stockholders' Agreement is ONE agreement; three texts of it must not diverge, even on a glyph.
SHA_26_TAIL = ("The Company shall obtain directors’ and officers’ liability insurance on or before the closing "
               "of the Seed Round and shall cause the Investor Representative to be covered under that policy "
               "in his capacities as Board observer or director and as Investor Representative. The Company "
               "represents that, as of the date hereof, it does not maintain such insurance.")

SHA_27 = ("2.7 No personal liability of the Investor Representative. The Investor Representative shall have no "
          "personal liability to any Investor for any action taken, or omitted to be taken, by him in good "
          "faith in his capacity as Investor Representative, except to the extent such action or omission "
          "constitutes fraud, bad faith, willful misconduct or gross negligence. The Investor Representative "
          "is not a fiduciary of, and owes no fiduciary duty to, any Investor, and may have interests that "
          "differ from those of any Investor. Each Investor, by executing this Agreement or a joinder to it, "
          "acknowledges and agrees to this Section 2.7. Nothing in this Section creates any liability or duty "
          "that would not otherwise exist, and nothing in this Section limits the indemnification provided "
          "under Section 2.6. This Section 2.7 survives termination of this Agreement.")

SL_5D = ("(d) Insurance. The Company represents that, as of the date hereof, it does not maintain directors' "
         "and officers' liability insurance. The Company shall obtain such insurance on or before the closing "
         "of the Seed Round and shall, upon obtaining it, cover the Investor Representative in both his "
         "capacity as Investor Representative and his capacity as a Board observer, and shall provide him "
         "with the policy's coverage limits and material terms upon request.")

SUB_IR = ("Investor Representative. The Investor acknowledges and agrees that, under the Side Letter "
          "Agreement, Ankur Jain serves as the Investor Representative and is authorized to give and receive "
          "notices, consents, waivers, elections and directions on behalf of the Investor, and that any such "
          "action binds the Investor, in each case except where that agreement requires the Investor's "
          "individual consent. The Investor further acknowledges that the Investor Representative may, in his "
          "sole discretion and at any time — without regard to any anniversary of the date hereof and "
          "without regard to the Company's achievement of any operating, financing or other milestone — "
          "require the Company to return undeployed proceeds of this offering to the investors on a pro rata "
          "basis in accordance with the Side Letter Agreement, and that the Company will indemnify the "
          "Investor Representative for good-faith actions taken in that capacity, including in respect of "
          "claims brought by an investor.")

CHANGES = {
 "Ankur_Jain_Strategic_Adviser_Agreement": [
   ("§1.1 role — drop approval limb",
    r"1\.1 Role\. Adviser serves as a non-exclusive Strategic Adviser.*?network\.",
    "1.1 Role. " + ADV_11),
   ("§1.2 approval → consultation",
    r"^1\.2 Financial materials — approval\..*?delayed\.$", ADV_12),
   ("§1.3 deemed approval → no condition / no remedy",
    r"^1\.3 Valuation, financing structure and investor terms\..*?for determination by the Board\.$", ADV_13),
   ("§2.5 de-hardcode the seed size",
    r"^2\.5 Anti-dilution — Seed Round only\..*?any other grant\.$", ADV_25),
   ("§2.7 83(b) inapplicable to vested shares",
    r"^2\.7 Mechanics; tax\..*?of each grant\.$", ADV_27),
   ("§4.3 D&O firm trigger at Seed close",
    r"^4\.3 D&O insurance\..*?Investor Representative\.$", ADV_43),
 ],
 "Cinderella_Corp_-_RSPA_-_Ankur_Jain": [
   ("header FMV note removed (inconsistency 1 of 3)",
    r"^Completed for Ankur Jain pursuant to the Strategic Adviser Agreement.*?confirm with counsel before execution\.$",
    "Completed for Ankur Jain pursuant to the Strategic Adviser Agreement dated ______, 2026. Share count "
    "reflects 4% of the Company’s fully-diluted capitalization as of the Effective Date, against "
    "3,000,000 shares of Class B Common Stock held by the Founder and a pre-Seed fully-diluted total of "
    "4,166,667 shares."),
   ("§1.1 price stated as Board-determined FMV (inconsistency 2 of 3)",
    r"1\.1 Shares\..*?\$0\.001 per Share\.", RSPA_11),
   ("§1.3 exemption → Reg D / 4(a)(2)",
    r"^1\.3 Securities exemption\..*?an available exemption\.$", RSPA_13),
   ("§2.4 milestone on objective achievement, not Company determination",
    r"^2\.4 Milestone grants\..*?governed exclusively by that separate agreement\.$", RSPA_24),
   ("Schedule A price (inconsistency 3 of 3)",
    r"Purchase price per Share: \$0\.001 \[FMV as determined by the Board\]",
    "Purchase price per Share: $0.001 (Board-determined fair market value as of the Effective Date)"),
 ],
 "Cinderella_Corp_-_Stockholders_Agreement_-_Ankur_Jain": [
   ("§2.6 D&O firm trigger at Seed close",
    r"The Company will use commercially reasonable efforts to include the Investor Representative under its "
    r"directors’ and officers’ liability insurance in his capacities as Board observer or director "
    r"and as Investor Representative\.", SHA_26_TAIL),
   ("§2.7 exculpation inserted", "INSERT_AFTER:2.6 Indemnification of the Investor Representative", SHA_27),
 ],
 "Cinderella_Corp_-_Side_Letter_Agreement_Revenue_Share_Final": [
   ("§5(d) D&O firm trigger at Seed close",
    r"^\(d\) Insurance\..*?coverage limits and material terms upon request\.$", SL_5D),
 ],
 "Cinderella_Corp_-_Subscription_Agreement_Final": [
   ("Investor Rep acknowledgment — capital return self-evident",
    r"^Investor Representative\. The Investor acknowledges and agrees that.*?claims brought by an investor\.$",
    SUB_IR),
 ],
}

# ───────────────────────── docx emission ─────────────────────────

SEC   = re.compile(r"^(\d{1,2})\.\s+[A-Z]")           # 1. SERVICES AND ROLE
SUB   = re.compile(r"^\d{1,2}\.\d{1,2}\s")            # 1.2 Financial presentation
LETTER= re.compile(r"^\(([a-z])\)\s")                 # (d) Insurance
BULLET= re.compile(r"^[••]\s*")

def style(doc):
    n = doc.styles["Normal"]
    n.font.name, n.font.size = "Calibri", Pt(10)
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.line_spacing = 1.08
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1.0)
        s.top_margin = s.bottom_margin = Inches(0.9)

def add_platform_table(doc):
    rows = [("Qualifying Platform","Meeting (25%)","Term sheet (25%)","Definitive (50%)","Total"),
            ("Netflix","1.25%","1.25%","2.50%","5.00%"),
            ("Apple","0.875%","0.875%","1.75%","3.50%"),
            ("Amazon Prime Video","0.875%","0.875%","1.75%","3.50%"),
            ("Disney / Hulu / ESPN","0.875%","0.875%","1.75%","3.50%"),
            ("All other Qualifying Platforms","0.25%","0.25%","0.50%","1.00%")]
    t = doc.add_table(rows=len(rows), cols=5)
    t.style = "Table Grid"
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j)
            c.text = ""
            run = c.paragraphs[0].add_run(v)
            run.font.size = Pt(9)
            if i == 0:
                run.bold = True
            c.paragraphs[0].paragraph_format.space_after = Pt(2)
    doc.add_paragraph()

def emit(name, paras, title):
    doc = Document(); style(doc)
    h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(title); r.bold = True; r.font.size = Pt(13)
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Revised 2026-10-07 — incorporating Adviser’s 2026-10-07 comments")
    sr.italic = True; sr.font.size = Pt(8.5); sr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    doc.add_paragraph()
    for p in paras:
        if p == TABLE_MARK:
            add_platform_table(doc); continue
        if p.strip() == title.strip() or p.strip() == title.upper().strip():
            continue
        if SEC.match(p) and len(p) < 90 and p == p.upper():
            par = doc.add_paragraph(); par.paragraph_format.space_before = Pt(10)
            run = par.add_run(p); run.bold = True; run.font.size = Pt(10.5)
        elif SEC.match(p):
            par = doc.add_paragraph(); par.paragraph_format.space_before = Pt(10)
            head, _, rest = p.partition(" ")
            num_end = p.find(". ", len(head)) if False else None
            m = re.match(r"^(\d{1,2}\.\s+[A-Z][A-Za-z &;,’/\-]*?)(?=\s[A-Z][a-z]|$)", p)
            if m:
                run = par.add_run(m.group(1)); run.bold = True
                par.add_run(p[len(m.group(1)):])
            else:
                run = par.add_run(p); run.bold = True
        elif SUB.match(p):
            par = doc.add_paragraph()
            m = re.match(r"^(\d{1,2}\.\d{1,2}\s+[^.]{0,80}?\.)", p)
            if m:
                run = par.add_run(m.group(1)); run.bold = True
                par.add_run(p[len(m.group(1)):])
            else:
                par.add_run(p)
        elif LETTER.match(p):
            par = doc.add_paragraph(); par.paragraph_format.left_indent = Inches(0.3)
            m = re.match(r"^(\(\w\)\s+[^.]{0,60}?\.)", p)
            if m:
                run = par.add_run(m.group(1)); run.bold = True
                par.add_run(p[len(m.group(1)):])
            else:
                par.add_run(p)
        elif BULLET.match(p):
            par = doc.add_paragraph(BULLET.sub("", p), style="List Bullet")
        else:
            doc.add_paragraph(p)
    out = OUT / f"{name}.docx"
    doc.save(out)
    return out

# ───────────────────────── run ─────────────────────────

TITLES = {
 "Ankur_Jain_Strategic_Adviser_Agreement": "STRATEGIC ADVISER AGREEMENT",
 "Cinderella_Corp_-_RSPA_-_Ankur_Jain": "RESTRICTED STOCK PURCHASE AGREEMENT",
 "Cinderella_Corp_-_Stockholders_Agreement_-_Ankur_Jain": "STOCKHOLDERS’ AGREEMENT",
 "Cinderella_Corp_-_Side_Letter_Agreement_Revenue_Share_Final": "SIDE LETTER AGREEMENT — REVENUE SHARE",
 "Cinderella_Corp_-_Subscription_Agreement_Final": "SUBSCRIPTION AGREEMENT",
}

failures, applied = [], []
for name, edits in CHANGES.items():
    text = load(name)
    text, had_table = excise_platform_table(text)
    paras = paragraphs(text)
    for label, find, repl in edits:
        if isinstance(find, str) and find.startswith("INSERT_AFTER:"):
            anchor = find.split(":", 1)[1]
            idx = next((i for i, p in enumerate(paras) if p.startswith(anchor)), None)
            if idx is None:
                failures.append(f"{name}: {label} — anchor not found"); continue
            paras.insert(idx + 1, repl); applied.append(f"{name}: {label}"); continue
        hit = False
        for i, p in enumerate(paras):
            if re.search(find, p, re.S):
                paras[i] = re.sub(find, lambda _m: repl, p, flags=re.S); hit = True; break
        if hit:
            applied.append(f"{name}: {label}")
        else:
            failures.append(f"{name}: {label} — source string not found")
    out = emit(name, paras, TITLES[name])
    print(f"  wrote {out.name}  ({out.stat().st_size:,} bytes, {len(paras)} paras"
          f"{', platform table rebuilt' if had_table else ''})")

# ⚠️ REGRESSION GUARD (Ankur, 2026-10-07): §2.2's heading, the definition of Qualifying Platform
# and §2.2(a) Platform tiers must all survive the table excision, and must sit in that order
# immediately before the rebuilt grid. They were all being swallowed.
_adv = Document(OUT / "Ankur_Jain_Strategic_Adviser_Agreement.docx")
_ps = [q.text.strip() for q in _adv.paragraphs]
_i = next((n for n, x in enumerate(_ps) if x.startswith("2.2 Distribution Milestone Equity")), None)
assert _i is not None, "REGRESSION: §2.2 heading missing from the Adviser Agreement"
assert "A \u201cQualifying Platform\u201d means Netflix" in _ps[_i], \
    "REGRESSION: the Qualifying Platform definition was excised with the table"
_a = next((n for n, x in enumerate(_ps) if x.startswith("(a) Platform tiers")), None)
assert _a is not None, "REGRESSION: §2.2(a) Platform tiers missing"
assert _a > _i, "REGRESSION: §2.2(a) must follow the §2.2 heading"
_b = next((n for n, x in enumerate(_ps) if x.startswith("(b) Milestones")), None)
assert _b is not None and _b > _a, "REGRESSION: §2.2(b) must follow §2.2(a)"
assert len(_adv.tables) == 1, f"expected exactly 1 platform table, found {len(_adv.tables)}"
assert [c.text for c in _adv.tables[0].rows[0].cells] == \
    ["Qualifying Platform", "Meeting (25%)", "Term sheet (25%)", "Definitive (50%)", "Total"], \
    "platform table header row is wrong"
print("  OK  \u00a72.2 heading \u2192 definition \u2192 \u00a72.2(a) \u2192 table \u2192 \u00a72.2(b) all present, in order")

print(f"\napplied {len(applied)} / {len(applied)+len(failures)} changes")
for a in applied:
    print("  OK  " + a)
if failures:
    print("\nFAILED:")
    for f in failures:
        print("  !!  " + f)
    sys.exit(1)
