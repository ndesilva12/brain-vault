# -*- coding: utf-8 -*-
"""Expense & Travel Policy — OCTOBER 2026 version. This is the agreed policy.

SOURCE: Making_Cinderella_Expense_Policy_October_2026.pdf, sent by Norman 2026-10-07 after Ankur
flagged that Side Letter Exhibit A carried a different, looser version. Transcribed verbatim.

⚠️ SUPERSEDES 2026-09-18-expense-policy-generator.py for Side Letter Exhibit A. The Sept 18
version is materially LOOSER and must not be attached to the Side Letter again:

  | Term                    | Sept 18 (do not use)                          | October (agreed) |
  |-------------------------|-----------------------------------------------|------------------|
  | Other-expense threshold | $5,000                                        | **$500**         |
  | Lodging cap             | $300, with an exception for NYC/SF/LA         | **$300 flat**    |
  | Approver                | "the Investor Representative"                 | **Ankur Jain, by email** |
  | SPVs                    | expressly carved out (§1.1 Entities not cov.) | **no carve-out** |
  | Compensation            | expressly carved out (§1.2)                   | **no carve-out** |
  | 90-day vendor grouping  | present                                       | absent           |
  | Documentation / 30-day  | absent                                        | **present (§6)** |

⚠️ TWO CONSEQUENCES OF THE DROPPED CARVE-OUTS, flagged for Norman in the review memo:
  1. §1 reaches "Cinderella Corp. or any entity funded by Cinderella Corp." With the Sept 18
     "Entities not covered" clause gone, that language arguably reaches the per-season SPVs, which
     run on a different cost base entirely.
  2. With the Sept 18 "Compensation excluded" clause gone, "payment or reimbursement" is arguable
     as reaching salary, not just expenses.
  Both are Norman's call to raise or accept; the policy is transcribed as agreed either way.
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path("/tmp/claude-0/-home-user-brain-vault/cdeb7b3c-9b38-5f86-b0fa-23ce6e519c4c/scratchpad"
           "/Expense_and_Travel_Policy_October_2026.docx")

# (kind, text) — kind: T title, S section heading, P prose, B bullet, b sub-bullet, A adopted line
POLICY = [
    ("T", "MAKING CINDERELLA"),
    ("T", "EXPENSE & TRAVEL POLICY"),
    ("S", "1.  Purpose and Scope"),
    ("P", "This policy applies to all founders, employees, contractors, advisers and other persons "
          "seeking payment or reimbursement from Cinderella Corp. or any entity funded by "
          "Cinderella Corp. Company-paid expenses must be reasonable, business-related, properly "
          "documented and consistent with this policy."),
    ("S", "2.  Air Travel"),
    ("B", "Until 100% of Seed investor capital has been returned to investors, all Company-paid "
          "flights must be booked in coach/economy class."),
    ("B", "After 100% of Seed investor capital has been returned:"),
    ("b", "Scheduled flight time of 4 hours or less: coach/economy."),
    ("b", "More than 4 hours and up to 7 hours: premium economy."),
    ("b", "More than 7 hours: business class."),
    ("B", "Travelers may use personal cash, points or miles to upgrade above the permitted class "
          "at no cost to the Company."),
    ("B", "Any exception to the permitted cabin class requires Ankur Jain’s prior written "
          "approval by email."),
    ("S", "3.  Hotels and Airbnb"),
    ("B", "Maximum lodging rate: $300 per night before taxes."),
    ("B", "Any lodging above the $300 nightly limit requires Ankur Jain’s prior written "
          "approval by email."),
    ("B", "Personal extensions, room upgrades and other incremental personal costs are not "
          "reimbursable."),
    ("S", "4.  Meals"),
    ("B", "Breakfast: up to $20 per day."),
    ("B", "Lunch: up to $30 per day."),
    ("B", "Dinner: up to $50 per day."),
    ("B", "Maximum daily meal allowance: $100."),
    ("B", "If a meal is provided by a conference, hotel, sponsor, school, investor or other third "
          "party, the corresponding meal allowance is not reimbursable."),
    ("B", "Alcohol and entertainment are not reimbursable unless they are part of a bona fide "
          "business-development expense and otherwise comply with this policy."),
    ("S", "5.  Other Business Expenses"),
    ("B", "Any single non-flight, non-lodging expense or commitment over $500 requires Ankur "
          "Jain’s prior written approval by email."),
    ("B", "This includes, without limitation, entertainment, consultants, vendors, subscriptions, "
          "equipment, event costs, gifts, marketing and other business-development expenditures."),
    ("B", "Reasonable ground transportation, parking, tolls and similar travel expenses are "
          "reimbursable when incurred for Company business."),
    ("S", "6.  Documentation and Reimbursement"),
    ("B", "Receipts or other reasonable supporting documentation are required for reimbursement."),
    ("B", "Expense reports should identify the business purpose and, for business-development "
          "meals or entertainment, the attendees."),
    ("B", "Reimbursement requests should be submitted within 30 days after the expense is "
          "incurred whenever practicable."),
    ("B", "Personal expenses are not reimbursable."),
    ("S", "7.  Exceptions and Amendments"),
    ("B", "No person may approve his or her own exception to this policy."),
    ("B", "Any exception requiring approval under this policy must be approved in advance by "
          "Ankur Jain via email."),
    ("B", "Until 100% of Seed investor capital has been returned, any material amendment to this "
          "policy requires the written approval of both the Chief Executive Officer and Ankur "
          "Jain in his capacity as Investor Representative."),
    ("A", "Adopted:  ____________________"),
]


def build(path=OUT):
    d = Document()
    for s in d.sections:
        s.top_margin = s.bottom_margin = Inches(1.0)
        s.left_margin = s.right_margin = Inches(1.0)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_after = Pt(9)
    st.paragraph_format.line_spacing = 1.06

    for kind, text in POLICY:
        p = d.add_paragraph()
        pf = p.paragraph_format
        if kind == "T":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.space_after = Pt(3)
            r = p.add_run(text); r.bold = True; r.font.size = Pt(12)
        elif kind == "S":
            pf.space_before, pf.space_after = Pt(13), Pt(6)
            r = p.add_run(text); r.bold = True; r.font.size = Pt(11)
        elif kind == "P":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.add_run(text)
        elif kind == "B":
            pf.left_indent, pf.space_after = Inches(0.28), Pt(5)
            p.add_run("•  " + text)
        elif kind == "b":
            pf.left_indent, pf.space_after = Inches(0.58), Pt(4)
            p.add_run("◦  " + text)
        elif kind == "A":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.space_before = Pt(22)
            p.add_run(text)

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    d.save(path)
    return path


if __name__ == "__main__":
    print("wrote", build())
