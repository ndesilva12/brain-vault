# -*- coding: utf-8 -*-
"""Render a .docx to PDF without LibreOffice.

⚠️ WHY THIS EXISTS. LibreOffice is installed in this container but cannot open ANY .docx —
`soffice --convert-to pdf` returns "Error: source file could not be loaded" even for a one-line
file written by python-docx, so it is the install that is broken, not our documents. Chromium is
present and its headless --print-to-pdf works, so the path is docx -> HTML -> PDF.

Fidelity is deliberately limited to what these contracts actually use: paragraphs, bold runs,
centre/justify alignment, left indents, and simple tables. That covers the full Cinderella legal
set. It is NOT a general-purpose converter — it ignores images, headers/footers, numbering fields
and styles beyond the above.
"""
import html
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

CHROME_CANDIDATES = [
    "/opt/pw-browsers/chromium-1194/chrome-linux/chrome",
    "/opt/pw-browsers/chromium/chrome-linux/chrome",
    "/usr/bin/chromium",
    "/usr/bin/google-chrome",
]

CSS = """
@page { size: Letter; margin: 1in; }
html { -webkit-print-color-adjust: exact; }
body { font-family: Calibri, Carlito, "DejaVu Sans", sans-serif; font-size: 10.5pt;
       line-height: 1.32; color: #000; margin: 0; }
p { margin: 0 0 7pt 0; orphans: 3; widows: 3; }
p.c { text-align: center; }
p.j { text-align: justify; }
b { font-weight: 700; }
table { border-collapse: collapse; width: 100%; margin: 8pt 0 10pt 0;
        font-size: 10pt; page-break-inside: avoid; }
td, th { border: 0.75pt solid #000; padding: 4pt 6pt; vertical-align: top; text-align: left; }
tr:first-child td { font-weight: 700; }
.pb { page-break-before: always; }
"""


def _runs_html(para):
    out = []
    for r in para.runs:
        t = html.escape(r.text).replace("\n", "<br>")
        if not t:
            continue
        if r.bold:
            t = f"<b>{t}</b>"
        if r.italic:
            t = f"<i>{t}</i>"
        out.append(t)
    return "".join(out) or "&nbsp;"


def _para_html(para, *, page_break=False):
    cls = []
    if page_break:
        cls.append("pb")
    al = para.alignment
    if al == WD_ALIGN_PARAGRAPH.CENTER:
        cls.append("c")
    elif al == WD_ALIGN_PARAGRAPH.JUSTIFY:
        cls.append("j")
    style = ""
    ind = para.paragraph_format.left_indent
    if ind:
        style = f' style="margin-left:{ind.inches:.2f}in"'
    c = f' class="{" ".join(cls)}"' if cls else ""
    return f"<p{c}{style}>{_runs_html(para)}</p>"


def _table_html(tbl):
    rows = []
    for row in tbl.rows:
        cells = "".join(f"<td>{html.escape(c.text)}</td>" for c in row.cells)
        rows.append(f"<tr>{cells}</tr>")
    return "<table>" + "".join(rows) + "</table>"


def to_html(docx_path, *, break_before=()):
    """break_before: paragraph-text prefixes that should start a new page (e.g. "EXHIBIT A")."""
    doc = Document(docx_path)
    body = doc.element.body
    # walk the body in document order so tables land between the right paragraphs
    para_by_el = {p._p: p for p in doc.paragraphs}
    tbl_by_el = {t._tbl: t for t in doc.tables}
    parts = []
    for el in body.iterchildren():
        if el in para_by_el:
            p = para_by_el[el]
            txt = p.text.strip()
            brk = any(txt.startswith(pfx) for pfx in break_before) if txt else False
            parts.append(_para_html(p, page_break=brk))
        elif el in tbl_by_el:
            parts.append(_table_html(tbl_by_el[el]))
    title = html.escape(Path(docx_path).stem)
    return (f"<!doctype html><html><head><meta charset='utf-8'><title>{title}</title>"
            f"<style>{CSS}</style></head><body>" + "".join(parts) + "</body></html>")


def convert(docx_path, pdf_path, *, break_before=(), workdir=None):
    chrome = next((c for c in CHROME_CANDIDATES if Path(c).exists()), None)
    assert chrome, f"no chromium found; tried {CHROME_CANDIDATES}"
    workdir = Path(workdir or Path(pdf_path).parent)
    workdir.mkdir(parents=True, exist_ok=True)
    tmp_html = workdir / (Path(pdf_path).stem + ".html")
    tmp_html.write_text(to_html(docx_path, break_before=break_before), encoding="utf-8")
    r = subprocess.run(
        [chrome, "--headless", "--no-sandbox", "--disable-gpu", "--no-pdf-header-footer",
         f"--print-to-pdf={pdf_path}", f"file://{tmp_html.resolve()}"],
        capture_output=True, text=True, timeout=300)
    assert Path(pdf_path).exists(), f"chromium produced no PDF:\n{r.stdout}\n{r.stderr}"
    return Path(pdf_path)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
    print("wrote", sys.argv[2])
