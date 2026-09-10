# -*- coding: utf-8 -*-
"""Sunjay Mathews Partner Agreement — DRAFT 4 (final).
DRAFT 3 plus: §2.6 Acceleration, Cause limbs (b)/(f)/(g) cut, conforming
cross-reference fixes in §2.3(c), §2.4 and §6.3, and the OpenAI typo."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = ("/tmp/claude-0/-home-user-brain-vault/cdeb7b3c-9b38-5f86-b0fa-23ce6e519c4c/"
       "scratchpad/Sunjay Mathews Partner Agreement DRAFT 4 (FINAL).docx")

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Inches(1.0)
    s.left_margin = s.right_margin = Inches(1.0)

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
st.font.color.rgb = RGBColor(0, 0, 0)
st.paragraph_format.space_after = Pt(10)
st.paragraph_format.line_spacing = 1.08


def p(text="", *, lead=None, indent=0.0, center=False, italic=False, bold=False,
      before=0, after=10, size=11, justify=True):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if indent:
        pf.left_indent = Inches(indent)
    if center:
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif justify:
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if lead:
        r = par.add_run(lead); r.bold = True; r.font.size = Pt(size)
    if text:
        r = par.add_run(text); r.italic = italic; r.bold = bold; r.font.size = Pt(size)
    return par


def head(text, size=12, before=16, after=8):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    r = par.add_run(text); r.bold = True; r.font.size = Pt(size)


def rule():
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(10)
    par.paragraph_format.space_after = Pt(4)
    r = par.add_run("_" * 92)
    r.font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
    r.font.size = Pt(8)


p("DRAFT — FOR DISCUSSION ONLY — NOT LEGAL ADVICE", bold=True, center=True,
  size=9.5, after=14, justify=False)

par = doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
par.paragraph_format.space_after = Pt(16)
r = par.add_run("PARTNER AGREEMENT — RESTRICTED STOCK, VESTING & MILESTONE EQUITY")
r.bold = True; r.font.size = Pt(13)

p("This Partner Agreement (this “Agreement”) is entered into as of [____], 2026 (the "
  "“Effective Date”) by and between Cinderella Corp., a Delaware corporation (the "
  "“Company”), and Sunjay Mathews (“Partner”).")

head("RECITALS", size=11.5, before=10, after=8)

p("The Making Cinderella concept, format, structure, narrative architecture, production "
  "methodology, format bible, and franchise architecture (collectively, the “Format”) "
  "were conceived, created, and developed solely by Norman C. de Silva.",
  lead="A.  Origination.  ")

p("Prior to the date on which Partner first became involved with the Company or the Format (the "
  "“Partner Start Date”, being April 28, 2026), Norman C. de Silva had, without any "
  "involvement or contribution by Partner: (i) created the Format and the associated format and "
  "development materials; (ii) originated and materially advanced the Company's relationship with "
  "Davidson College, which by the Partner Start Date had progressed to an advanced stage / the "
  "point of a signed letter of intent; (iii) originated the Company's celebrity-attachment "
  "strategy and its school-procurement approach; and (iv) established the Company's capital, "
  "production, and intellectual-property structures.",
  lead="B.  Prior development.  ")

p("Partner acknowledges and agrees that Recitals A and B are accurate; that he made no material "
  "contribution to the original conception or creation of the concept; and that he joined an "
  "existing venture with existing assets, relationships, and momentum.",
  lead="C.  Acknowledgment.  ")

p("These Recitals are contractual statements of fact relied upon by the Company in entering into "
  "this Agreement and in granting the equity described herein. They are not mere preamble and are "
  "incorporated into this Agreement as though set out in full in the body.",
  lead="D.  Reliance.  ")

rule()
head("1.  TITLE, ROLE, AUTHORITY AND PUBLIC REPRESENTATION")

p("Partner's title is “Founding Partner” (the “Title”).", lead="1.1  Title.  ")

p("Partner is not a founder, co-founder, or originator of the Company or of the Format. Partner "
  "shall not, during the term or at any time thereafter, hold himself out as — or knowingly "
  "permit any third party to describe him as — a founder, co-founder, originator, creator, "
  "co-creator, inventor, or architect of the Company, the Format, or Making Cinderella, and shall "
  "promptly correct any such description within his knowledge and control. The Company and the "
  "Format were created by and remain attributable to Norman C. de Silva.",
  lead="1.2  Not a founder; no origination claim.  ")

p("The Title is granted for the term only, is tied to Partner's role, and may be changed by the "
  "Company on written notice if Partner's role or scope changes. Upon termination for any reason, "
  "Partner's right to use the Title ceases immediately, and Partner shall not thereafter use the "
  "Title or any variant — including any formulation using “former,” “founding,” "
  "“co-founder,” “originally,” or “of Making Cinderella” — or any other "
  "formulation implying origination, co-equal standing, or continuing affiliation.",
  lead="1.3  Title is licensed, not owned.  ")

p("Partner will lead partnerships, media, legal coordination, and commercial execution, reporting "
  "to the Chief Executive Officer, and will devote full-time or substantially all business time "
  "to the Company.", lead="1.4  Role.  ")

p("For the avoidance of doubt, Partner:", lead="1.5  No special rights or authority.  ")
for t in [
    "(a)  holds Class A Common Stock (1 vote per share) only — Partner does not hold, and is not "
    "entitled to, Class B Common Stock or any super-voting stock (held solely by Norman C. de "
    "Silva);",
    "(b)  has no board seat, no board-observer right, and no right to designate any director;",
    "(c)  has no veto, consent, approval, protective, information, or anti-dilution right of any "
    "kind beyond those held by holders of Class A Common Stock generally;",
    "(d)  has no authority to bind the Company or to execute agreements on its behalf except as "
    "expressly authorized in writing by the CEO for a specific matter — this includes no "
    "authority to sign or deliver any letter of intent, term sheet, memorandum of understanding, "
    "side letter, or definitive agreement, and no authority to make any oral or written "
    "commitment as to price, economics, exclusivity, timing, or attachment;",
    "(e)  is subject to the same transfer restrictions, right of first refusal, co-sale, and "
    "drag-along provisions as all other Class A holders under the Stockholders' Agreement; and",
    "(f)  is not an officer of the Company. The Title is a business title only, does not "
    "constitute appointment to any office named in the Company's Bylaws, and confers no officer "
    "authority, signing authority, or fiduciary office.",
]:
    p(t, indent=0.35)

p("", lead="1.6  Sole spokesperson; public representation.", after=6)
for lead_txt, body in [
    ("(a)  Sole spokesperson.  ",
     "Norman C. de Silva — or such person as the CEO designates in writing for a specific matter "
     "— is the sole authorized spokesperson for the Company."),
    ("(b)  Major public representation.  ",
     "Partner shall not, without the CEO's prior written approval, represent the Company in any "
     "press interview, podcast, broadcast, panel, or conference appearance, or issue any press "
     "statement, concerning the Company, the Format, any school, any talent, any sponsor, any "
     "capital partner, or any distributor."),
    ("(c)  No statement of Company position.  ",
     "Partner shall not state, imply, or negotiate a Company position — as to economics, "
     "structure, ownership, timing, exclusivity, or attachment — in any meeting or communication "
     "without the CEO's prior approval of that position."),
    ("(d)  No representation of status.  ",
     "Partner shall not represent to any third party that any school, talent, sponsor, capital "
     "partner, or distributor is attached, committed, engaged, or closed unless a definitive "
     "agreement has been executed. (This mirrors the Company's standing rule that no party may "
     "represent attachment before signature.)"),
    ("(e)  Marks.  ",
     "Partner shall not use the Company's names, logos, or marks in any personal account, "
     "venture, or content."),
    ("(f)  Survival; injunctive relief.  ",
     "This Section 1.6 and Sections 1.2–1.3 survive termination. Partner acknowledges that a "
     "breach would cause irreparable harm for which monetary damages are inadequate, and that the "
     "Company is entitled to injunctive relief without posting bond, in addition to all other "
     "remedies."),
]:
    p(body, lead=lead_txt, indent=0.35)

rule()
head("2.  BASE EQUITY — 20%")

p("The Company grants Partner restricted Class A Common Stock equal to twenty percent (20%) of "
  "the Company (the “Base Shares”), measured on the fully-diluted capitalization of the "
  "Company immediately PRIOR to the Seed Round (i.e., pre-money / pre-Seed).",
  lead="2.1  Grant.  ")

p("Because the Base Shares are measured pre-Seed, Partner is diluted by the Seed Round on the "
  "same basis as Norman and Greg Kristof. For illustration only: on a pre-Seed capitalization of "
  "Norman 72% / Partner 20% / Ankur Jain 4% / Greg Kristof 4%, and assuming a [15]% Seed Round in "
  "which Ankur Jain's position does not dilute (per the Company's separate arrangement with him), "
  "Norman, Partner, and Greg collectively absorb more than 15% dilution to hold Ankur flat and "
  "still deliver 15% to the Seed investors. On that assumption, Partner's approximate post-Seed "
  "position would be ~16.9% (not the ~17% a simple pro-rata calculation would suggest). This "
  "figure is illustrative only and must be confirmed against the actual Seed Round terms and cap "
  "table model before use. Partner has no anti-dilution protection with respect to the Seed Round "
  "or any subsequent issuance.",
  lead="2.2  Seed dilution applies.  ")

p("The Base Shares vest as follows:", lead="2.3  Vesting.  ", after=6)
p("(a)  Fifty percent (50%) of the Base Shares (being ten percent (10%) of the Company as "
  "measured under Section 2.1) are fully vested on the Effective Date; and", indent=0.35)
p("(b)  the remaining fifty percent (50%) of the Base Shares vest in twenty-four (24) equal "
  "monthly installments, with the first installment vesting on the last day of the first full "
  "calendar month following the Effective Date and each subsequent installment vesting on the "
  "last day of each full calendar month thereafter, in each case subject to Partner's continuous "
  "service through the applicable vesting date.", indent=0.35)
p("(c)  No installment vests in respect of a partial calendar month. Where the Company terminates "
  "Partner's service without Cause (as defined in Section 6.2), Partner's service shall be deemed "
  "to continue through the end of any notice period for purposes of this Section 2.3.",
  indent=0.35)

p("Except as provided in Section 2.6, on termination of service for any reason, all unvested Base "
  "Shares are immediately forfeited to (or repurchasable by) the Company at the lower of cost or "
  "fair market value. Vested shares remain subject to the Stockholders' Agreement.",
  lead="2.4  Forfeiture.  ")

p("The Base Shares are subject to ordinary, pro-rata dilution on all issuances (subject to "
  "Section 2.2 as to the Seed Round). Partner has no anti-dilution protection.",
  lead="2.5  Dilution.  ")

p("", lead="2.6  Acceleration.", after=6)
for lead_txt, body in [
    ("(a)  Termination without Cause.  ",
     "If the Company terminates Partner's service without Cause, the next six (6) monthly "
     "installments under Section 2.3(b) that would otherwise have vested shall accelerate and "
     "vest in full as of the termination date."),
    ("(b)  Change of Control.  ",
     "Immediately prior to the closing of a Change of Control, all then-unvested Base Shares "
     "shall accelerate and vest in full, subject to Partner's continuous service through such "
     "closing."),
    ("(c)  Milestone Shares unaffected.  ",
     "For the avoidance of doubt, this Section 2.6 applies only to the Base Shares. Milestone "
     "Shares are performance-based and do not accelerate; any Track not fully earned as of the "
     "relevant date is forfeited in accordance with Section 3. Milestone Shares for a Track "
     "properly earned before the closing of a Change of Control but not yet issued shall be "
     "issued at that closing."),
    ("(d)  Change of Control defined.  ",
     "“Change of Control” means (i) a merger, consolidation, share exchange or "
     "reorganization of the Company after which more than fifty percent (50%) of the combined "
     "voting power of the surviving entity is held by persons who were not stockholders of the "
     "Company immediately prior; (ii) a sale of capital stock by stockholders (other than in a "
     "public offering), in one transaction or a series of related transactions, after which more "
     "than fifty percent (50%) of the combined voting power of the Company is held by persons who "
     "were not stockholders immediately prior; or (iii) a sale, transfer or other disposition of "
     "all or substantially all of the assets of the Company. For the avoidance of doubt, the "
     "grant of a license of the Format to any special purpose vehicle, production partner or "
     "distributor in the ordinary course of the Company's business does not constitute a Change "
     "of Control."),
    ("(e)  Termination shortly before a Change of Control.  ",
     "If the Company terminates Partner's service without Cause within the six (6) months "
     "preceding the closing of a Change of Control, Section 2.6(b) applies as though Partner's "
     "service had continued through such closing."),
]:
    p(body, lead=lead_txt, indent=0.35)

rule()
head("3.  MILESTONE EQUITY — UP TO AN ADDITIONAL 5% (TO A 25% MAXIMUM)")

p("Partner may earn up to an additional five percent (5%) of the Company in Class A Common Stock "
  "(the “Milestone Shares”), bringing Partner's maximum equity to twenty-five percent "
  "(25%), by achieving Track A and/or Track B below. Each Track, if achieved, earns two and "
  "one-half percent (2.5%), and the Tracks stack — achieving both earns 5% in the aggregate. Each "
  "Track may be earned once only; a second qualifying event under a Track already earned earns "
  "nothing further.", lead="3.1  The Milestone Award.  ")

p("Each Track must be fully earned on or before December 31, 2029 (the “Outside Date”). "
  "Any Track not fully earned by the Outside Date is permanently forfeited.",
  lead="3.2  Outside Date.  ")

p("Each Track is binary. Partial or preliminary performance earns nothing. The Tracks are "
  "independent of one another: achievement, failure, or clawback of one Track does not affect "
  "Milestone Shares properly earned under the other.",
  lead="3.3  All-or-nothing; Tracks independent.  ")

head("TRACK A — Lead or Presenting Sponsor, Any Season", size=11, before=12, after=6)
p("The Milestone is earned if:", after=6)
for lead_txt, body in [
    ("(a)  Sponsor.  ",
     "The sponsor is one of the following brands: Nike, Adidas, Gatorade, Coca-Cola, Pepsi, "
     "Apple, Microsoft, Google, SoFi, Chase, State Farm, Allstate, AT&T, Verizon, Marriott, "
     "Chevrolet, Ford, Tesla, Anthropic, OpenAI, SpaceX (illustrative examples only), or any "
     "other brand approved by the Board, in its sole discretion."),
    ("(b)  Status.  ",
     "The sponsor holds Lead Sponsor or Presenting Sponsor status for any single season of any "
     "Company school project, or the Board otherwise determines, in its sole discretion, that the "
     "sponsor is the most prominent sponsor of that season. The sponsorship need not cover more "
     "than one school or more than one season, and no minimum dollar amount is required."),
    ("(c)  Definitive agreement.  ",
     "A definitive, fully executed sponsorship agreement is in place (a term sheet, LOI, or "
     "non-binding commitment does not qualify)."),
    ("(d)  Procuring cause.  ",
     "Partner is the direct procuring cause of the sponsorship, as determined by the Board in its "
     "sole discretion. Deals sourced through Norman, existing Company relationships, retained "
     "agencies, or third-party finders do not qualify."),
    ("(e)  Arm's length.  ",
     "The sponsor is not an affiliate of, and has no material pre-existing economic relationship "
     "with, Partner or his affiliates."),
    ("(f)  Sustained performance.  ",
     "The agreement remains in effect, and has not been terminated or materially reduced other "
     "than due to Company breach, for twelve (12) months following execution. Clawback: if it is "
     "terminated or materially reduced within that period, the Milestone Shares attributable to "
     "Track A are forfeited (or repurchasable at cost), whether or not previously issued. "
     "Forfeiture under this paragraph does not affect any Milestone Shares earned under Track B."),
]:
    p(body, lead=lead_txt, indent=0.35)

head("TRACK B — Board-Approved Celebrity Attachment, Season 2 or Later", size=11, before=12,
     after=6)
p("The Milestone is earned if:", after=6)
p("A celebrity is attached, on a definitive, fully executed basis, as talent for any Company "
  "school project, for Season 2 or any later season. A celebrity procured for Season 1 does not "
  "qualify, regardless of when the Board's determination under (b) is made.",
  lead="(a)  Attachment.  ", indent=0.35)
p("The Board determines, in its sole discretion, that Partner directly procured that attachment "
  "and that the attachment is properly attributable to Partner.",
  lead="(b)  Procuring cause and attribution.  ", indent=0.35)
p("Board approval under this Track B is not a formality and may be withheld or granted for any "
  "reason, in the Board's sole discretion.")

p("Whether a Track has been achieved is determined by the Board in good faith, whose "
  "determination is final and binding. Partner will provide such documentation as the Board "
  "reasonably requests.", lead="3.4  Determination.  ")

p("With respect to each Track separately, Partner must be in continuous service with the Company, "
  "and not in material breach of this Agreement, on the date that Track is achieved and on the "
  "date the corresponding Milestone Shares are issued.", lead="3.5  Service condition.  ")

p("The Milestone Shares for each Track are issued as restricted Class A Common Stock upon "
  "achievement of that Track, measured on the fully-diluted capitalization as of that Track's "
  "achievement date, subject to ordinary dilution thereafter, with no anti-dilution protection.",
  lead="3.6  Issuance, measurement and dilution.  ")

rule()
head("4.  TAX")

p("Partner is strongly advised to consult his own tax adviser and to consider filing an 83(b) "
  "election within thirty (30) days of each grant. The Company makes no tax representation.",
  lead="4.1  83(b).  ")
p("The 10% vesting on the Effective Date under Section 2.3(a) is immediately vested and may "
  "generate immediate taxable income to Partner equal to the fair market value of those shares.",
  lead="4.2  Immediate vesting — tax note.  ")
p("The Company may withhold, or require payment of, applicable taxes as a condition to issuance.",
  lead="4.3  Withholding.  ")

rule()
head("5.  CONFIDENTIALITY; INTELLECTUAL PROPERTY; NON-CIRCUMVENTION")

p("Partner will hold all non-public Company information in confidence and use it solely for the "
  "Company's benefit.", lead="5.1  Confidentiality.  ")

p("All work product, materials, concepts, formats, and deliverables Partner creates in connection "
  "with the Company are works made for hire and/or assigned to the Company. Partner acknowledges "
  "that the Making Cinderella format, brand, and franchise IP are and remain the exclusive "
  "property of the Company, and Partner asserts no authorship, creation, or ownership claim to "
  "them.", lead="5.2  IP assignment.  ")

p("Consistent with the Recitals, Partner:", lead="5.3  Waiver of creation and credit claims.  ",
  after=6)
for t in [
    "(a)  irrevocably waives any claim to authorship, co-authorship, creation, or co-creation of "
    "the Format, and any claim that he originated, co-originated, or materially contributed to "
    "the conception of the Format or of the Company;",
    "(b)  irrevocably waives any claim to a “Created by,” “Developed by,” "
    "“Format by,” or equivalent credit on any production, and acknowledges the Company's "
    "standing requirement that “Created by Norman C. de Silva” appear in first position "
    "on every project;",
    "(c)  waives all moral rights and rights of paternity, integrity, and attribution in the "
    "Format and any work product, to the fullest extent waivable under applicable law; and",
    "(d)  agrees that any producer, executive producer, or other on-screen credit he may receive "
    "is a matter of the Company's sole discretion, is not a credit of creation or authorship, and "
    "creates no ownership, backend, or format right.",
]:
    p(t, indent=0.35)
p("This Section 5.3 survives termination and applies whether or not any equity vests.")

p("", lead="5.4  Non-competition; non-solicitation.", after=6)
for lead_txt, body in [
    ("(a)  Restricted Project.  ",
     "“Restricted Project” means any business, venture, production or project that "
     "(i) pairs one or more celebrities, public figures, athletes or prominent alumni with one or "
     "more collegiate athletic programs for the purpose of financing, capitalizing or materially "
     "supporting that program and producing, licensing or distributing documentary, docuseries or "
     "similar audiovisual content concerning it, or (ii) is otherwise substantially similar to "
     "the Format or to Making Cinderella."),
    ("(b)  Restriction.  ",
     "During the term and for twenty-four (24) months thereafter, Partner shall not, directly or "
     "indirectly, whether as founder, owner, employee, officer, director, consultant, adviser, "
     "investor, lender, producer or in any other capacity: (i) develop, launch, operate or "
     "participate in any Restricted Project; (ii) assist, advise or enable any third party to do "
     "so; or (iii) lend his name, likeness or endorsement to any Restricted Project. Passive "
     "ownership of less than two percent (2%) of a publicly traded company is permitted."),
    ("(c)  Relationships expressly unrestricted.  ",
     "Nothing in this Agreement restricts Partner from initiating, continuing or maintaining any "
     "personal or professional relationship with any person or entity, whenever and however that "
     "relationship arose — including any school, athlete, celebrity, talent representative, "
     "sponsor, capital partner, distributor or production partner introduced to Partner through "
     "or in connection with the Company — during the term or at any time thereafter. Partner's "
     "obligations under Sections 5.1 and 5.2 continue to apply to Company confidential "
     "information and intellectual property in any such communication."),
    ("(d)  Non-solicitation of personnel.  ",
     "During the term and for twelve (12) months thereafter, Partner shall not solicit for "
     "employment or engagement any employee, officer or contractor of the Company."),
]:
    p(body, lead=lead_txt, indent=0.35)

p("Partner's existing minority sports-club interests are disclosed and permitted, provided they "
  "do not conflict with Partner's obligations.", lead="5.5  Outside activities.  ")

rule()
head("6.  TERM AND TERMINATION")

p("Partner's service may be terminated by Partner at any time on 30 days' written notice, and by "
  "the Company (i) without Cause on 30 days' written notice or (ii) for Cause, immediately upon "
  "written notice and without any cure period.", lead="6.1  Termination.  ")

p("“Cause” means any of the following, as determined by the Board in its good-faith "
  "judgment:", lead="6.2  Cause.  ", after=6)
for t in [
    "(a)  any breach of this Agreement, including Sections 1.2, 1.5, 1.6, 5.1, 5.2, 5.3 or 5.4;",
    "(b)  any act or omission involving dishonesty, fraud, misappropriation, breach of fiduciary "
    "duty, or material lack of candor toward the Company or the CEO;",
    "(c)  indictment for, conviction of, or plea of guilty or no contest to, any felony or any "
    "crime involving dishonesty or moral turpitude; or",
    "(d)  any conduct, statement, publication, endorsement, or public association — whether or "
    "not occurring in the course of Partner's service, whether or not lawful, and regardless of "
    "the views expressed — that in the Board's good-faith judgment is reasonably likely to cause "
    "material harm to the Company's reputation or goodwill, or to impair or jeopardize the "
    "Company's actual or prospective relationships with any school, athletic conference, "
    "governing body, athlete, talent, talent representative, sponsor, capital partner, "
    "distributor, or production partner.",
]:
    p(t, indent=0.35)
p("The Board's determination that Cause exists is final and binding absent manifest bad faith. "
  "Cause may be found whether or not the conduct was intentional and whether or not any actual "
  "harm has yet occurred.")

p("On termination: unvested Base Shares are forfeited (Section 2.4, subject to Section 2.6); "
  "unearned Milestone Shares are permanently forfeited; vested shares are retained subject to the "
  "Stockholders' Agreement and any repurchase rights. Sections 1.2, 1.3, 1.6, 4–5 and 7 survive.",
  lead="6.3  Effect.  ")

rule()
head("7.  GENERAL")
p("Delaware governing law; entire agreement together with the Restricted Stock Purchase Agreement "
  "and Stockholders' Agreement; amendments in writing signed by both parties; Partner may not "
  "assign; counterparts and electronic signature permitted.")

rule()
p("COMPANY: Cinderella Corp.", bold=True, before=10, after=16, justify=False)
p("By:  ______________________________", after=2, justify=False)
p("Name:  Norman C. de Silva", after=2, justify=False)
p("Title:  Founder & Chief Executive Officer", after=2, justify=False)
p("Date:  ______________", after=20, justify=False)

p("PARTNER:  ______________________________", bold=True, after=2, justify=False)
p("Sunjay Mathews", after=2, justify=False)
p("Date:  ______________", justify=False)

doc.save(OUT)
print("wrote", OUT)
