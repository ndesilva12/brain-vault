# -*- coding: utf-8 -*-
"""Subscription Agreement — v2.
Base: Drive > Legal > Seed Round > Final > "Cinderella Corp - Subscription Agreement Final"
(file 1vYhWtma6R--_GNOwBUShcPWU7iaQ52yh). Loeb's text verbatim; additions respond to Ankur's
2026-09-18 email items 1, 2, 4 and 9."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = ("/tmp/claude-0/-home-user-brain-vault/cdeb7b3c-9b38-5f86-b0fa-23ce6e519c4c/scratchpad/"
       "Cinderella Corp - Subscription Agreement Final (v2).docx")
d = Document()
for s in d.sections:
    s.top_margin = s.bottom_margin = Inches(1.0); s.left_margin = s.right_margin = Inches(1.0)
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
st.font.color.rgb = RGBColor(0,0,0); st.paragraph_format.space_after = Pt(10)
st.paragraph_format.line_spacing = 1.08

def p(t="", **k):
    par=d.add_paragraph(); pf=par.paragraph_format
    pf.space_after=Pt(k.get("after",10)); pf.space_before=Pt(k.get("before",0))
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if k.get("center") else WD_ALIGN_PARAGRAPH.JUSTIFY
    if k.get("indent"): pf.left_indent=Inches(k["indent"])
    r=par.add_run(t); r.bold=k.get("bold",False); r.italic=k.get("italic",False)
    if k.get("size"): r.font.size=Pt(k["size"])
def num(n, title, body):
    par=d.add_paragraph(); pf=par.paragraph_format
    pf.space_after=Pt(10); pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    r=par.add_run(f"{n}. {title}"); r.bold=True
    par.add_run("  "+body)

p("SUBSCRIPTION AGREEMENT", bold=True, center=True, size=13, after=14)
p("The undersigned individual or entity (the “Investor”) hereby subscribes for shares of Class A "
  "Common Stock, par value $0.001 (the “Shares”) of Cinderella Corp., a Delaware corporation (the "
  "“Company”), at the price of $[____] per Share, on the terms and conditions set forth in this "
  "Subscription Agreement (this “Subscription Agreement”), the Side Letter Agreement (as defined "
  "below), and the Certificate of Incorporation, as amended, of the Company, which is attached "
  "hereto as Exhibit A.")
p("The Investor hereby irrevocably subscribes for that number of Shares set forth on the signature "
  "page of this Subscription Agreement.")
p("Side Letter Agreement. Concurrently with the execution of this Subscription Agreement, and as a "
  "condition to the Investor’s purchase of the Shares, the Investor and the Company are entering "
  "into a Side Letter Agreement among the Company and the investors party thereto (the “Side "
  "Letter Agreement”), pursuant to which the Investor is entitled to certain preferential payments "
  "based on the gross revenue of the Company. The Investor acknowledges and agrees that (a) the "
  "Side Letter Agreement forms part of the same transaction as, and is a material inducement to, "
  "the Investor’s subscription hereunder; (b) the preferential payments provided for in the Side "
  "Letter Agreement are in addition to, and are not credited or offset against, any dividend, "
  "distribution or other amount payable to the Investor in respect of the Shares; and (c) the "
  "Investor’s rights in respect of the Shares, including the right to receive dividends and "
  "distributions pro rata with other holders of capital stock of the Company, are unaffected by "
  "the Side Letter Agreement.")

# NEW — Ankur items 1 and 2
p("Investor Representative. The Investor acknowledges and agrees that, under the Side Letter "
  "Agreement, Ankur Jain serves as the Investor Representative and is authorized to give and "
  "receive notices, consents, waivers, elections and directions on behalf of the Investor, and "
  "that any such action binds the Investor, in each case except where that agreement requires the "
  "Investor’s individual consent. The Investor further acknowledges that the Investor "
  "Representative may, in his sole discretion and at any time, require the Company to return "
  "undeployed proceeds of this offering to the investors on a pro rata basis in accordance with "
  "the Side Letter Agreement, and that the Company will indemnify the Investor Representative for "
  "good-faith actions taken in that capacity, including in respect of claims brought by an "
  "investor.")

# NEW — Ankur items 4 and 9
p("Governing Corporate Documents. The Shares are shares of Class A Common Stock of the Company, "
  "carrying one (1) vote per share. The Company has two classes of common stock outstanding, Class "
  "A Common Stock and Class B Common Stock, as set forth in the Certificate of Incorporation "
  "attached as Exhibit A, and no class or series of preferred stock is authorized or outstanding. "
  "The Shares are additionally subject to the Company’s Stockholders’ Agreement, a copy of which "
  "has been furnished to the Investor and which contains transfer restrictions, a right of first "
  "refusal, co-sale rights, drag-along obligations and related provisions. The Investor is a party "
  "to the Stockholders’ Agreement or shall execute a joinder to it concurrently with this "
  "Subscription Agreement. In the event of any conflict between this Subscription Agreement and "
  "the Side Letter Agreement as to the preferential payments or the return of undeployed capital, "
  "the Side Letter Agreement controls.")

p("The Investor hereby represents and warrants to the Company as follows:", before=4)
REPS = [
 ("Formation; Authority.", "The Investor, if a juridicial entity, has been duly formed and is validly existing in the state of its formation, with all requisite power and authority to enter into this Subscription Agreement, to carry out the provisions and conditions hereof, and to consummate the transactions contemplated hereby. The Investor, if an individual, is at least 21 years old and is legally competent."),
 ("Binding Agreement.", "This Subscription Agreement is a legal, valid and binding agreement of the Investor, enforceable against the Investor in accordance with its terms, subject to applicable bankruptcy, insolvency, fraudulent conveyance, reorganization, moratorium and similar laws affecting creditors’ rights and remedies generally and subject, as to enforceability, to general principles of equity, including principles of commercial reasonableness, good faith and fair dealing (regardless of whether enforcement is sought in a proceeding at law or in equity)."),
 ("Speculative Nature; Risk of Loss.", "The Investor recognizes that the Company is a highly speculative venture, involving a high degree of financial risk, and has read and understands the risk factors set forth on Exhibit B hereto. The Investor acknowledges that the risk factors set forth on Exhibit B do not purport to be a complete statement of all risks associated with an investment in the Company."),
 ("No Registration of Shares.", "The Investor recognizes that the Shares and the proposed sale of the Shares to the Investor have not been and will not be registered under the Securities Act of 1933, as amended (the “Act”), and that, therefore, the Shares may not be sold or otherwise transferred by the Investor unless the Shares are subsequently registered under the Act or unless in the opinion of counsel for the Company, a sale, assignment or transfer of the Shares may be made without registration thereunder."),
 ("Ability to Bear Economic Risk.", "The Investor (i) understands that he, she or it must bear the economic risk of an investment in the Shares for an indefinite period of time and (ii) is able to bear such economic risk, including the total loss of his, her or its investment."),
 ("Financial Knowledge and Advice.", "The Investor has such knowledge and experience in financial affairs that he, she or it is capable of evaluating, or has employed the services of an investment advisor, attorney or accountant to evaluate, on his, her or its behalf, the merits and risks of purchasing the Shares."),
 ("No Other Representations or Warranties.", "Except as set forth herein and in the Side Letter Agreement, no representations or warranties have been made to the Investor by the Company or any agent, employee or affiliate of the Company, and that in entering into this transaction the Investor is not relying on any information other than the results of any independent investigation by or on behalf of the Investor."),
 ("Acquisition for Investment Purposes.", "The Investor confirms that the Investor is acquiring the Shares subscribed for herein solely for the Investor’s own account, for investment purposes, and not with a view to the distribution or resale of such Shares."),
 ("Access to Information.", "The Investor acknowledges having been furnished with all relevant terms and conditions of this investment, including this Subscription Agreement, the Side Letter Agreement, the Stockholders’ Agreement, the Certificate of Incorporation attached as Exhibit A, and the risk factors attached as Exhibit B, and such other documents, materials and information as the Investor (and Investor’s purchaser representative, if any) deems necessary or appropriate for evaluating an investment in the Company. The Investor confirms that the Investor (and Investor’s purchaser representative, if any) has read and understands these materials and has made such further investigation of the Company as was deemed appropriate to obtain additional information to verify the accuracy of such materials and to evaluate the merits and risks of this investment. The Investor acknowledges that the Investor (and Investor’s purchaser representative, if any) has had the opportunity to ask questions of, and receive answers from, the Company concerning the terms and conditions of the offering and the information contained in the offering materials."),
 ("Securities Laws Requirements.", "The Investor recognizes that the securities laws and regulations of certain states, including the state of which the Investor is a resident, may impose additional requirements relating to this offering and Investor’s purchase of the Shares in the Company. The Investor hereby agrees to execute and to comply with the terms of any supplements or amendments to this Subscription Agreement which are required by the Company."),
 ("No Governmental Endorsement.", "The Investor understands that no federal or state agency has recommended or endorsed the purchase of the Shares as an investment or passed on the adequacy of the information set forth in any of the other offering materials."),
 ("No General Solicitation or Advertising.", "The Investor acknowledges that neither the Company nor any person acting on behalf of the Company offered to sell, or sold to the Investor, the Shares by means of any form of general solicitation or general advertising."),
 ("Legal and Financial Advisors.", "The Investor acknowledges that the Investor has been advised to consult with Investor’s own legal counsel and financial advisors regarding legal matters concerning the Company and to consult with Investor’s tax advisor regarding the tax consequences of participating in the Company."),
 ("No Special Purpose Entity.", "If the Investor is a corporation, Company, association, trust, unincorporated organization or other entity, the Investor represents that it (or, if it is a wholly-owned subsidiary, its parent corporation) has not been formed for the specific purpose of making an investment in the Company and that the Investor has the full power and authority under its governing instruments to execute this Subscription Agreement on behalf of the Investor and that the Investor has the full power and authority under such instruments to become a Stockholder in the Company."),
 ("Ownership of Investor.", "If the Investor is a corporation, Company, association, trust, unincorporated organization or other entity, the Investor represents that any equity or other owners of the Investor share in all the gains or losses of all investments of the Company in the same way and on the basis of their proportional ownership and do not have non-proportionate or non-pro rata Shares in specified investments of the Investor. Based on most recent valuations available: (i) less than twenty-five percent (25%) of Investor’s assets are owned by “benefit plan investors” as defined in regulations of the United States Department of Labor concerning those categories of assets that constitute assets of an employee benefit plan, and (ii) the Investor agrees to notify the Company promptly if the percentage of its assets owned by benefit plan investors should equal or exceed twenty-five percent (25%)."),
 ("Irrevocable Subscription.", "The Investor hereby agrees that this subscription is irrevocable and that the representations and warranties set forth in this Subscription Agreement shall survive the acceptance hereof by the Company and the subsequent purchase of the Shares by the Investor."),
 ("Effectiveness of Agreements and Representations.", "The agreements and representations herein set forth shall become effective and binding upon the Investor, the Investor’s legal representatives, heirs, successors and assigns, upon the Company’s acceptance of the Investor’s subscription."),
 ("Reaffirmations.", "Any representation or warranty made hereunder will be deemed to be reaffirmed at any time the Investor makes an additional investment in the Company. The act of making such additional investment will be evidence of such reaffirmation."),
 ("Lock-Up.", "If requested in writing by the managing underwriter of an initial public offering of the Company, he, she or it will not sell, transfer or otherwise dispose of any securities of the Company for such period of time after completion of such initial public offering (not to exceed one (1) year) as the managing underwriter may request in writing, provided such period of time is no longer than the shortest period of time required of any other purchaser of Shares in the Company."),
 ("Accredited Investor.", "The Investor is an “accredited investor” as defined under Rule 501 of Regulation D of the Act."),
]
for i,(t,b) in enumerate(REPS,1): num(i,t,b)

p("The Investor shall promptly notify the Company when any representation made herein is no longer "
  "accurate.")
p("Neither this Subscription Agreement nor any provisions hereof shall be waived, modified, "
  "discharged or terminated except by an instrument in writing signed by the party against whom "
  "any waiver, modification, discharge or termination is sought. This Subscription Agreement shall "
  "be binding upon and inure to the benefit of the parties and their successors and permitted "
  "assigns. If the Investor is more than one person, the obligation of the Investor shall be joint "
  "and several, and the agreements, presentations, warranties and acknowledgments herein contained "
  "shall be deemed to be made by and be binding upon each such person and its successors and "
  "assigns.")
p("This Subscription Agreement, the Side Letter Agreement, the Stockholders’ Agreement, and the "
  "other agreements or documents referred to herein or therein contain the entire agreement of the "
  "parties, and there are no representations, covenants or other agreements except as stated or "
  "referred to herein or therein. For the avoidance of doubt, nothing in this Subscription "
  "Agreement limits, supersedes or modifies any right of the Investor under the Side Letter "
  "Agreement.")
p("This Subscription Agreement is irrevocable and not transferable or assignable by the Investor "
  "without the Company’s prior written consent.")
p("This Subscription Agreement shall be governed by and construed in accordance with the internal "
  "laws of the State of Delaware without regard to conflicts of laws principles.")
p("Any term or provision of this Subscription Agreement that is invalid or unenforceable in any "
  "jurisdiction shall, as to that jurisdiction, be ineffective to the extent of such invalidity or "
  "unenforceability without rendering invalid or unenforceable the remaining terms or provisions "
  "of this Subscription Agreement or affecting the validity or enforceability of any of the terms "
  "or provisions of this Subscription Agreement in any other jurisdiction.")
p("This Subscription Agreement may be executed in one or more counterparts, each of which may be "
  "delivered electronically and shall constitute an original, and all of which when taken together "
  "shall constitute one and the same instrument.")
p("BY EXECUTING THE SIGNATURE PAGE TO THIS SUBSCRIPTION AGREEMENT, THE INVESTOR AGREES TO BE BOUND "
  "BY THE FOREGOING.", bold=True)

d.add_page_break()
p("IN WITNESS WHEREOF, the undersigned has executed this Subscription Agreement as of the [____] "
  "day of ____________, 20[__].", after=18)
for lab in ["Print Name of Investor","Signature of Investor",
            "Tax Identification No. or Social Security No.",
            "Address (Post Office Box Not Acceptable)","Mailing Address (If Different)"]:
    par=d.add_paragraph(); par.paragraph_format.space_after=Pt(2); par.add_run("_"*52)
    q=d.add_paragraph(); q.paragraph_format.space_after=Pt(12); q.add_run(lab)
p("Subscription Amount: $______________________", after=4)
p("Number of Shares: _________________________", after=16)
p("ACCEPTANCE", bold=True, after=6)
p("The undersigned, on behalf of the Company, hereby executes this Subscription Agreement and "
  "accepts the foregoing subscription this [____] day of ____________, 20[__].", after=14)
p("CINDERELLA CORP", after=2)
par=d.add_paragraph(); par.paragraph_format.space_after=Pt(2); par.add_run("_"*52)
for lab in ["By:","Name: Norman de Silva","Title: President & Founder"]:
    q=d.add_paragraph(); q.paragraph_format.space_after=Pt(0); q.add_run(lab)

d.add_page_break()
p("EXHIBIT A", bold=True, center=True, size=12, after=4)
p("CERTIFICATE OF INCORPORATION", bold=True, center=True, size=11, after=10)
p("[Attach the Amended and Restated Certificate of Incorporation of Cinderella Corp., as amended "
  "to date, authorizing 6,000,000 shares of Class A Common Stock and 4,000,000 shares of Class B "
  "Common Stock and no preferred stock.]", italic=True, center=True)

d.add_page_break()
p("EXHIBIT B", bold=True, center=True, size=12, after=4)
p("RISK FACTORS", bold=True, center=True, size=11, after=10)
p("An investment in the Shares involves a high degree of risk. The Investor should carefully "
  "consider the following risks, together with all other information contained in the offering "
  "materials, before deciding to invest. The risks described below are not exhaustive, and "
  "additional risks not presently known to the Company, or that the Company currently deems "
  "immaterial, may also impair its business and the value of the Shares.")
RISKS = [
 ("No Operating History; Early Stage.","The Company was recently formed, has a limited operating history, and has not yet generated meaningful revenue. It has not completed a season of programming and there is no basis on which to evaluate its ability to execute its business plan."),
 ("Total Loss of Investment.","The Investor may lose the entire amount of its investment. The Company may never become profitable, may never make any distribution, and may cease operations."),
 ("Illiquidity; No Public Market.","The Shares are not registered under the Act and are subject to substantial transfer restrictions, including those in the Stockholders’ Agreement. No public market exists or is expected to develop, and the Investor may be unable to sell the Shares at any price for an indefinite period."),
 ("No Binding Talent Commitments.","The Company’s business model depends on attaching prominent celebrities to participating institutions. As of the date hereof, the Company has no definitive, fully executed talent agreement. Discussions with talent and their representatives are preliminary and may not result in any binding commitment."),
 ("Non-Binding School Commitments.","Letters of intent executed with academic institutions are non-binding expressions of interest. Institutions may withdraw, and definitive agreements may never be executed on acceptable terms, or at all."),
 ("No Committed Distribution Partner.","The Company has no executed agreement with any streaming platform, network or other distributor. There is no assurance that any programming will be licensed, produced or distributed, or that any license fee will be obtained on acceptable terms."),
 ("No Committed Project Capital.","Each school-level special purpose vehicle is expected to require substantial capital. The Company has no binding commitment for such capital, and discussions with prospective capital partners may not result in funding."),
 ("Regulatory Risk — NCAA and Name, Image and Likeness.","The regulation of collegiate athlete compensation is evolving rapidly, including through implementation of the House settlement, NCAA rulemaking, conference policy, and federal and state legislation. Changes in these rules, or their interpretation or enforcement, could materially impair or prohibit aspects of the Company’s model, including its arrangements with institutions and athletes."),
 ("Institutional Control and Compliance.","The Company’s arrangements are structured to preserve institutional control over athletics programs. A determination by the NCAA, an institution, a conference or a regulator that any arrangement is impermissible could require restructuring, result in penalties to a participating institution, or terminate a project."),
 ("Dependence on Key Personnel.","The Company depends substantially on Norman de Silva and a small number of other individuals. The loss of any of them, or the inability to attract additional personnel, would materially harm the Company."),
 ("Concentration.","The Company expects to operate a small number of projects at any time. Poor competitive performance, an injury, a coaching change, adverse publicity or the failure of a single project could disproportionately affect the Company’s results."),
 ("Preferential Payments Reduce Available Cash.","Under the Side Letter Agreement, a substantial percentage of the Company’s gross revenue is payable to investors before the Company applies revenue to its operating expenses. This will reduce cash otherwise available to fund operations and growth and may require the Company to seek additional financing."),
 ("Return of Undeployed Capital.","Under the Side Letter Agreement, the Investor Representative may, in his sole discretion and at any time, require the Company to return undeployed proceeds of this offering to investors. An exercise of that right would reduce the capital available to the Company to execute its business plan, may require the Company to curtail or abandon planned activities, and may occur at a time when the Company is unable to replace the returned capital on acceptable terms or at all."),
 ("Control by the Class B Stockholder.","The holder of Class B Common Stock holds ten (10) votes per share and controls the Company. Holders of Class A Common Stock will have limited ability to influence corporate decisions, including the composition of the board, the declaration of dividends, and whether to pursue or accept a sale of the Company."),
 ("Reliance on the Investor Representative.","The Investor Representative is authorized to act on behalf of all investors under the Side Letter Agreement, and his actions bind the Investor. The Investor Representative is not the Investor’s agent or fiduciary, may have interests that differ from the Investor’s, and is indemnified by the Company for good-faith actions taken in that capacity."),
 ("Dilution.","The Company expects to issue additional equity in future financings, to service providers, and under any equity incentive plan. Such issuances will dilute the Investor’s percentage ownership."),
 ("Intellectual Property.","The Company’s value depends significantly on its format and franchise intellectual property. Formats and concepts receive limited protection under applicable law, competitors may develop similar programming, and the Company may lack the resources to enforce its rights."),
 ("Reliance on Third Parties.","The Company depends on production companies, talent agencies, sponsors, institutions and capital partners, none of which it controls. Disputes with, or the non-performance of, any of these parties could materially harm the Company."),
 ("Conflicts of Interest.","Officers and affiliates of the Company may hold interests in, and receive compensation from, project-level entities and other affiliates. These relationships may present conflicts between their interests and those of the Investor."),
 ("Competition.","The Company competes for programming, talent, capital and institutional relationships with well-established production companies, media platforms, athlete collectives and investment firms, many of which have far greater resources."),
 ("Tax Consequences.","The tax consequences of an investment in the Shares are complex and depend on the Investor’s circumstances. The Company makes no representation as to any tax treatment, and the Investor should consult its own tax adviser."),
 ("Additional Financing.","The Company may require additional capital sooner than anticipated. Such capital may be unavailable, or available only on terms that are dilutive or otherwise adverse to existing holders."),
]
for i,(t,b) in enumerate(RISKS,1): num(i,t,b)

d.save(OUT); print("wrote", OUT)
